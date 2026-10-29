import os
from pathlib import Path
import openpyxl

from ops.pulse.kit.build_pulse_kit import build_pulse_excel


def test_build_pulse_excel_private_tracker():
    """Verifica que el tracker privado para Datalaria se genera con todas las hojas y datos."""
    dest_path = build_pulse_excel(is_template=False)
    assert dest_path.exists()
    assert dest_path.name == "Pulse_Private_Tracker.xlsx"
    
    wb = openpyxl.load_workbook(dest_path, data_only=False)
    expected_sheets = ["Instrucciones", "Activos", "Kanban", "Métricas", "Dashboard"]
    for s in expected_sheets:
        assert s in wb.sheetnames
        
    ws_act = wb["Activos"]
    # Debe contener cabecera + al menos 135 activos catalogados
    assert ws_act.max_row >= 136
    
    # Comprobar columnas clave
    assert ws_act.cell(row=1, column=1).value == "ID"
    assert ws_act.cell(row=1, column=10).value == "Paridad"
    assert ws_act.cell(row=1, column=15).value == "Score SEO"
    assert ws_act.cell(row=1, column=20).value == "Visitas 30d"


def test_build_pulse_excel_template():
    """Verifica que la plantilla del Kit comercial se genera limpiamente."""
    dest_path = build_pulse_excel(is_template=True)
    assert dest_path.exists()
    assert dest_path.name == "Pulse_Content_Ops_Template.xlsx"
    
    wb = openpyxl.load_workbook(dest_path, data_only=False)
    assert "Activos" in wb.sheetnames
    assert "Dashboard" in wb.sheetnames
