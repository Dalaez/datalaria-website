"""
ops/pulse/tests/test_kit_packaging.py
-------------------------------------
Suite de pruebas automatizadas para el Kit comercial descargable de Datalaria Pulse.
Especificación: Hito 5.6 (specs/apps/pulse/spec.md).

Valida:
1. Estructura, fórmulas e integridad de la plantilla Excel y taxonomía.
2. Migrador de bases de datos de Notion (normalización de estados y fechas).
3. Escáner universal de repositorios de contenidos sin dependencias de terceros.
4. Empaquetado e integridad del archivo ZIP final para Lemon Squeezy.
"""

import sys
import zipfile
from pathlib import Path
import openpyxl
import pytest

ROOT_DIR = Path(__file__).resolve().parent.parent.parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from ops.pulse.kit.build_pulse_kit import build_pulse_excel
from ops.pulse.kit.package_kit import package_kit, ZIP_OUTPUT_PATH
from ops.pulse.migrate_notion import migrate_notion, normalize_status, parse_date_str
from packages.pulse.pulse_repo_scanner import scan_directory, audit_seo, parse_simple_front_matter


def test_template_excel_structure(tmp_path):
    """Verifica que la plantilla del Kit comercial contiene las 6 hojas y los 5 ejemplos guiados."""
    excel_path = build_pulse_excel(is_template=True)
    assert excel_path.exists()
    assert excel_path.name == "Pulse_Content_Ops_Template.xlsx"

    wb = openpyxl.load_workbook(excel_path, data_only=False)
    expected_sheets = ["Instrucciones", "Taxonomía", "Activos", "Kanban", "Métricas", "Dashboard"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Hoja faltante: {s}"

    # Comprobar hoja Taxonomía
    ws_tax = wb["Taxonomía"]
    assert ws_tax.cell(row=1, column=1).value == "Tipo de Activo"
    assert ws_tax.cell(row=2, column=1).value == "POST"
    assert ws_tax.cell(row=1, column=2).value == "Estado"
    assert ws_tax.cell(row=2, column=2).value == "Idea"

    # Comprobar hoja Activos: 1 cabecera + 5 ejemplos guiados
    ws_act = wb["Activos"]
    assert ws_act.max_row == 6
    assert ws_act.cell(row=2, column=2).value == "POST"
    assert ws_act.cell(row=2, column=3).value == "Published"
    assert ws_act.cell(row=3, column=3).value == "Scheduled"
    assert ws_act.cell(row=4, column=3).value == "Draft"
    assert ws_act.cell(row=5, column=3).value == "Idea"
    assert ws_act.cell(row=6, column=2).value == "KIT"

    # Comprobar fórmulas universales en Dashboard (INV-PULSE-5)
    ws_dash = wb["Dashboard"]
    # Total Activos
    assert "COUNTA(" in str(ws_dash.cell(row=4, column=4).value)
    # Artículos Publicados
    assert 'COUNTIF(Activos!C2:C5, "Published")' in str(ws_dash.cell(row=5, column=4).value) or "COUNTIF" in str(ws_dash.cell(row=5, column=4).value)


def test_migrate_notion_csv(tmp_path):
    """Verifica la importación y normalización determinista de un CSV exportado de Notion."""
    csv_file = tmp_path / "notion_export.csv"
    csv_content = (
        "Name,Status,Date,Category,Tags,Summary\n"
        "Guía de Optimización,En redacción,2026-11-05,Operaciones,\"ops,lean\",Borrador inicial\n"
        "Nuevo Producto B2B,Live,15/10/2026,Estrategia,\"c-level,b2b\",Lanzamiento suite\n"
        "Idea de Podcast,To Do,,Marketing,\"audio\",Explorar podcast semanal\n"
        "Post sobre Machine Learning,Calendarizado,2026-12-01,IA,\"ml,python\",En cola de revisión\n"
    )
    csv_file.write_text(csv_content, encoding="utf-8")

    out_yaml = tmp_path / "backlog_imported.yaml"
    assets = migrate_notion(csv_file, output_path=out_yaml, out_format="yaml")

    assert len(assets) == 4
    assert out_yaml.exists()

    # Mapeo de estados normalizados
    by_title = {a["title_es"]: a for a in assets}
    assert by_title["Guía de Optimización"]["status"] == "draft"
    assert by_title["Guía de Optimización"]["publish_date"] == "2026-11-05"
    assert by_title["Nuevo Producto B2B"]["status"] == "published"
    assert by_title["Nuevo Producto B2B"]["publish_date"] == "2026-10-15"
    assert by_title["Idea de Podcast"]["status"] == "idea"
    assert by_title["Idea de Podcast"]["publish_date"] is None
    assert by_title["Post sobre Machine Learning"]["status"] == "scheduled"
    assert by_title["Post sobre Machine Learning"]["publish_date"] == "2026-12-01"


def test_pulse_repo_scanner(tmp_path):
    """Verifica que el escáner universal detecta front matter, paridad y calcula score SEO."""
    content_dir = tmp_path / "content"
    es_dir = content_dir / "es" / "posts" / "articulo-test"
    en_dir = content_dir / "en" / "posts" / "articulo-test"
    es_dir.mkdir(parents=True, exist_ok=True)
    en_dir.mkdir(parents=True, exist_ok=True)

    # Post ES
    (es_dir / "index.md").write_text(
        "---\n"
        "title: 'Guía Cuantitativa de Prueba para Comités de Dirección'\n"
        "date: '2026-10-20'\n"
        "draft: false\n"
        "categories: ['Estrategia']\n"
        "tags: ['c-level', 'finanzas', 'analitica']\n"
        "summary: 'Este es un resumen estructurado para validar la auditoría de SEO y lectura.'\n"
        "---\n\n"
        "# Introducción\n\n"
        "Contenido de prueba con suficiente longitud para validar el escáner. " * 30 +
        "\n\n[Enlace Interno](/es/apps/) y una imagen ![Grafico](/img/chart.png)\n",
        encoding="utf-8"
    )

    # Post EN (paridad completa)
    (en_dir / "index.md").write_text(
        "---\n"
        "title: 'Quantitative Strategy Guide for Executive Committees'\n"
        "date: '2026-10-20'\n"
        "draft: false\n"
        "---\n\n"
        "# Introduction\n\n"
        "English test content to ensure complete language parity across directories.\n",
        encoding="utf-8"
    )

    results = scan_directory(content_dir)
    assert len(results) == 1
    post = results[0]

    assert post["ID"] == "articulo-test"
    assert post["Paridad"] == "complete"
    assert post["Título (ES)"] == "Guía Cuantitativa de Prueba para Comités de Dirección"
    assert post["Título (EN)"] == "Quantitative Strategy Guide for Executive Committees"
    assert post["Palabras ES"] > 100
    assert post["Score SEO"] >= 60
    assert post["Estado"] == "Published"


def test_package_kit_zip():
    """Verifica que el archivo ZIP comercial se genera correctamente con todos los ficheros."""
    zip_path = package_kit()
    assert zip_path.exists()
    assert zip_path == ZIP_OUTPUT_PATH
    assert zip_path.stat().st_size > 10000 # Al menos 10 KB

    with zipfile.ZipFile(zip_path, "r") as zf:
        bad_file = zf.testzip()
        assert bad_file is None, f"Archivo ZIP corrupto: {bad_file}"
        names = zf.namelist()
        expected_files = [
            "Pulse_Content_Ops_Template.xlsx",
            "pulse_repo_scanner.py",
            "migrate_notion.py",
            "notion_power_query.m",
            "GUIA_PULSE_KIT.md",
            "LICENSE.txt"
        ]
        for ef in expected_files:
            assert ef in names, f"Fichero faltante en el ZIP comercial: {ef}"
