import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('packages/Suite_02_Toma_Decisiones/03_Business_Case_VAN_TIR/[ES]_Business_Case_VAN_TIR/Business_Case_VAN_TIR_Datalaria_ES.xlsx', data_only=False)
ws_sens = wb['Sensibilidad & Tornado']
ws_assump = wb['Supuestos & Drivers']

print('=== Supuestos & Drivers: D50:E57 ===')
for r in range(50, 58):
    print(f'Fila {r}: B={ws_assump.cell(r, 2).value}, C(Base)={ws_assump.cell(r, 3).value}, D(Low)={ws_assump.cell(r, 4).value}, E(High)={ws_assump.cell(r, 5).value}')

print('\n=== Supuestos & Drivers: D34:D44 (Base values) ===')
for r in range(34, 45):
    print(f'Fila {r}: B={ws_assump.cell(r, 2).value}, D={ws_assump.cell(r, 4).value}')

print('\n=== Sensibilidad & Tornado: Filas 16-17 (op_low, op_high) ===')
for r in [16, 17]:
    for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
        print(f'{col}{r} = {ws_sens[f"{col}{r}"].value}')

print('\n=== Sensibilidad & Tornado: Filas 18-19 (gv_low, gv_high) ===')
for r in [18, 19]:
    for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']:
        print(f'{col}{r} = {ws_sens[f"{col}{r}"].value}')
