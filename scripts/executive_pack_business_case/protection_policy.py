#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
protection_policy.py
====================
Módulo transversal de política de protección de celdas y seguridad OpenXML
para el Executive Decision Pack Business Case (Datalaria.com).

Regla de oro determinista:
- Celda CON fórmula (comienza por '=' o ArrayFormula) -> locked=True, fondo #F1F5F9 (salvo decorativas).
- Celda SIN fórmula (entradas, textos, números, vacías) -> locked=False, fondo #FFFFFF (salvo decorativas).
- Protección de hoja: selectUnlockedCells=False, selectLockedCells=False (atributos son restricciones en XML).
- Contraseña oficial: "Datalaria2026"
"""

from openpyxl.styles import Protection, PatternFill
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.formula import ArrayFormula

PASSWORD_PROTECT = "Datalaria2026"
FILL_INPUT = PatternFill(fill_type="solid", fgColor="FFFFFF")    # Editable (#FFFFFF)
FILL_FORMULA = PatternFill(fill_type="solid", fgColor="F1F5F9")  # Protegida (#F1F5F9)


def is_formula(cell):
    """Determina si una celda contiene una fórmula de cálculo."""
    v = cell.value
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def enforce_protection_policy(ws, decorative_cells=frozenset(), extra_rows=0, extra_cols=0):
    """
    Pasada final y exhaustiva sobre todo el rango usado (+ margen opcional).
    iter_rows con límites explícitos recorre todas las celdas (incluso vacías),
    garantizando que ninguna celda sin fórmula quede bloqueada por omisión.
    """
    max_row = ws.max_row + extra_rows
    max_col = ws.max_column + extra_cols
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if is_formula(cell):
                cell.protection = Protection(locked=True, hidden=False)
                if cell.coordinate not in decorative_cells:
                    cell.fill = FILL_FORMULA
            else:
                cell.protection = Protection(locked=False, hidden=False)
                if cell.coordinate not in decorative_cells:
                    cell.fill = FILL_INPUT


def apply_sheet_protection(ws, allow_structure=True):
    """
    En CT_SheetProtection (ECMA-376 OpenXML):
    - sheet/objects/scenarios = True (blindado)
    - Los atributos selectUnlockedCells y selectLockedCells representan restricciones:
      False (0 en XML) = permitido al usuario (permite interactuar y editar celdas desbloqueadas).
    """
    ws.protection.set_password(PASSWORD_PROTECT)
    ws.protection.sheet = True
    ws.protection.objects = True
    ws.protection.scenarios = True
    ws.protection.selectUnlockedCells = False      # Permite seleccionar y editar celdas desbloqueadas
    ws.protection.selectLockedCells = False        # Permite seleccionar celdas bloqueadas (lectura/copia)
    if allow_structure:
        ws.protection.insertRows = False
        ws.protection.deleteRows = False
        ws.protection.formatCells = False
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.protection.sort = False
        ws.protection.autoFilter = False


def finalize_sheet(ws, decorative_cells=frozenset(), extra_rows=0, extra_cols=0):
    """Aplica la política de protección y activa la protección de la hoja como última operación."""
    enforce_protection_policy(ws, decorative_cells, extra_rows, extra_cols)
    apply_sheet_protection(ws)
