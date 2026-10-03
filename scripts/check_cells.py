import zipfile
import re
import openpyxl

wb = openpyxl.load_workbook('packages/ES_Matriz_McKinsey_GE_Datalaria.xlsx')
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- Sheet: {sheetname} ---")
    print(f"Protection enabled: {ws.protection.sheet}")
    print(f"selectUnlockedCells: {ws.protection.selectUnlockedCells}")
    print(f"selectLockedCells: {ws.protection.selectLockedCells}")
    unlocked = []
    locked = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                if cell.protection and cell.protection.locked is False:
                    unlocked.append(cell.coordinate)
                else:
                    locked.append(cell.coordinate)
    print(f"Unlocked count: {len(unlocked)} -> {unlocked[:15]}")
    print(f"Locked count: {len(locked)} -> {locked[:5]}")
