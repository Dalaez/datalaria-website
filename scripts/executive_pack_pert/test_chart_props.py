import openpyxl
from openpyxl.chart import AreaChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Test"

ws["A1"] = "Días"
ws["B1"] = "Densidad f(x)"

import numpy as np
mu = 192.5
sigma = 10.14
for i in range(25):
    x_val = round(mu + (-3.5 + i * (7.0/24.0)) * sigma, 1)
    # normal density
    y_val = (1.0 / (sigma * np.sqrt(2 * np.pi))) * np.exp(-0.5 * ((x_val - mu)/sigma)**2)
    ws.cell(row=i+2, column=1, value=x_val)
    ws.cell(row=i+2, column=2, value=round(y_val, 4))

chart = AreaChart()
chart.title = "Distribución de Probabilidad del Cronograma (Curva de Campana)"
chart.style = 13
chart.width = 18
chart.height = 11

chart.x_axis.title = "Duración en Días Hábiles"
chart.x_axis.axPos = "b"
chart.x_axis.tickLblPos = "nextTo"
chart.x_axis.majorTickMark = "out"
chart.x_axis.majorGridlines = ChartLines()
chart.x_axis.tickLblSkip = 3  # Show every 3rd label so it's clean and readable!

chart.y_axis.title = "Densidad de Probabilidad f(x)"
chart.y_axis.axPos = "l"
chart.y_axis.tickLblPos = "nextTo"
chart.y_axis.majorGridlines = ChartLines()

data = Reference(ws, min_col=2, min_row=1, max_row=26)
cats = Reference(ws, min_col=1, min_row=2, max_row=26)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

ws.add_chart(chart, "D2")
wb.save("scripts/executive_pack_pert/test_output_chart.xlsx")
print("Saved test_output_chart.xlsx")
