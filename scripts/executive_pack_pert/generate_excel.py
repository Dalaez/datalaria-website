#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_excel.py
=================
Genera los modelos de cálculo en Excel (.xlsx) de alta dirección para el Executive Decision Pack:
1. packages/Suite_03_Control_Operativo/03_Estimacion_PERT_3_Puntos/[ES]_Estimacion_PERT_3_Puntos/Estimacion_PERT_Estocastica_ES.xlsx
2. packages/Suite_03_Control_Operativo/03_Estimacion_PERT_3_Puntos/[EN]_3Point_PERT_Estimation/Stochastic_PERT_Estimation_EN.xlsx

Estructura de 4 Pestañas Funcionales:
- Pestaña 1: "Dashboard Ejecutivo" / "Executive Dashboard"
- Pestaña 2: "Estimación WBS 3 Puntos" / "WBS 3-Point Estimation"
- Pestaña 3: "Camino Crítico & Análisis Z" / "Critical Path & Z-Score Analysis"
- Pestaña 4: "Buffers & Plan de Aceleración" / "Schedule Buffers & Crashing"

Estándar de Seguridad y Compatibilidad OpenXML ECMA-376:
- Celdas de entrada (usuario): locked=False, fondo blanco #FFFFFF.
- Celdas de fórmulas y títulos: locked=True, fondo #F1F5F9 o corporativo.
- Contraseña oficial: "Datalaria2026"
- ECMA-376: selectUnlockedCells = False, selectLockedCells = False
- Fórmulas estándar compatibles con Microsoft Excel y Google Sheets (NORMDIST, NORMINV, SQRT, SUMIF).
"""

import os
import sys
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.chart import AreaChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import NumRef, NumData, NumVal, StrRef, StrData, StrVal, AxDataSource, NumDataSource

from protection_policy import finalize_protection

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# --- PALETA CORPORATIVA DATALARIA ---
COLOR_NAVY_DARK = "0F172A"       # Slate 900
COLOR_NAVY_MED = "1E293B"        # Slate 800
COLOR_NAVY_LIGHT = "334155"      # Slate 700
COLOR_BLUE_ACCENT = "2563EB"     # Blue 600
COLOR_BLUE_LIGHT = "EFF6FF"      # Blue 50
COLOR_CYAN_ACCENT = "0284C7"     # Sky 600

COLOR_GREEN_BG = "D1FAE5"        # Emerald 100
COLOR_GREEN_TEXT = "065F46"      # Emerald 800
COLOR_GREEN_BORDER = "10B981"    # Emerald 500

COLOR_AMBER_BG = "FEF3C7"        # Amber 100
COLOR_AMBER_TEXT = "92400E"      # Amber 800
COLOR_AMBER_BORDER = "F59E0B"    # Amber 500

COLOR_RED_BG = "FEE2E2"          # Red 100
COLOR_RED_TEXT = "991B1B"        # Red 800
COLOR_RED_BORDER = "DC2626"      # Red 600

COLOR_BG_CARD = "F8FAFC"         # Slate 50
COLOR_BG_INPUT = "FFFFFF"        # Blanco puro para inputs
COLOR_BG_FORMULA = "F1F5F9"      # Slate 100 para fórmulas
COLOR_BORDER_LIGHT = "CBD5E1"    # Slate 300
COLOR_BORDER_DARK = "64748B"     # Slate 500

FONT_NAME = "Segoe UI"

# Estilos reutilizables
font_title = Font(name=FONT_NAME, size=15, bold=True, color="FFFFFF")
font_subtitle = Font(name=FONT_NAME, size=10, italic=True, color="94A3B8")
font_section = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY_DARK)
font_tbl_header = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
font_card_title = Font(name=FONT_NAME, size=8.5, bold=True, color="64748B")
font_card_val = Font(name=FONT_NAME, size=16, bold=True, color=COLOR_NAVY_DARK)
font_card_sub = Font(name=FONT_NAME, size=8, italic=True, color="64748B")
font_body = Font(name=FONT_NAME, size=9.5, color="1E293B")
font_body_bold = Font(name=FONT_NAME, size=9.5, bold=True, color="1E293B")
font_input = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_BLUE_ACCENT)

fill_dark = PatternFill(fill_type="solid", fgColor=COLOR_NAVY_DARK)
fill_med = PatternFill(fill_type="solid", fgColor=COLOR_NAVY_MED)
fill_blue_accent = PatternFill(fill_type="solid", fgColor=COLOR_BLUE_ACCENT)
fill_blue_card = PatternFill(fill_type="solid", fgColor=COLOR_BLUE_LIGHT)
fill_card = PatternFill(fill_type="solid", fgColor=COLOR_BG_CARD)
fill_formula = PatternFill(fill_type="solid", fgColor=COLOR_BG_FORMULA)
fill_input = PatternFill(fill_type="solid", fgColor=COLOR_BG_INPUT)

fill_green = PatternFill(fill_type="solid", fgColor=COLOR_GREEN_BG)
fill_amber = PatternFill(fill_type="solid", fgColor=COLOR_AMBER_BG)
fill_red = PatternFill(fill_type="solid", fgColor=COLOR_RED_BG)

thin_border_side = Side(style="thin", color=COLOR_BORDER_LIGHT)
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_thick_side = Side(style="medium", color=COLOR_BLUE_ACCENT)
border_card = Border(left=border_thick_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)


def build_wbs_data(lang='ES'):
    """Genera 25 paquetes de trabajo realistas para el proyecto de transformación digital."""
    if lang == 'ES':
        tasks = [
            ("1.1", "Fase 1: Arquitectura & Diseño", "Definición de Arquitectura Cloud & Microservicios", "SÍ", 6.0, 10.0, 20.0),
            ("1.2", "Fase 1: Arquitectura & Diseño", "Modelado Conceptual de Datos & Seguridad IAM", "NO", 4.0, 7.0, 14.0),
            ("1.3", "Fase 1: Arquitectura & Diseño", "Especificación de APIs & Contratos OpenAPI", "SÍ", 5.0, 8.0, 17.0),
            ("1.4", "Fase 1: Arquitectura & Diseño", "Revisión & Aprobación en Architecture Board", "SÍ", 3.0, 5.0, 12.0),
            ("1.5", "Fase 1: Arquitectura & Diseño", "Aprovisionamiento de Infraestructura IaC Terraform", "NO", 4.0, 6.0, 12.0),

            ("2.1", "Fase 2: Desarrollo Core", "Desarrollo del Motor Transaccional de Conciliación", "SÍ", 12.0, 18.0, 32.0),
            ("2.2", "Fase 2: Desarrollo Core", "Implementación de Autenticación OAuth2 / OIDC", "NO", 5.0, 8.0, 15.0),
            ("2.3", "Fase 2: Desarrollo Core", "Módulo de Procesamiento de Pagos & Pasarelas", "SÍ", 10.0, 16.0, 28.0),
            ("2.4", "Fase 2: Desarrollo Core", "Motor de Reglas de Negocio & Enrutamiento", "SÍ", 8.0, 14.0, 24.0),
            ("2.5", "Fase 2: Desarrollo Core", "Sistema de Auditoría & Trazabilidad de Logs", "NO", 4.0, 7.0, 13.0),

            ("3.1", "Fase 3: Integración & APIs", "Conector Bidireccional con ERP Core (SAP/Oracle)", "SÍ", 10.0, 17.0, 30.0),
            ("3.2", "Fase 3: Integración & APIs", "Integración con Pasarelas Bancarias & Adquirentes", "SÍ", 8.0, 14.0, 26.0),
            ("3.3", "Fase 3: Integración & APIs", "Conector con CRM Salesforce & Servicio al Cliente", "NO", 5.0, 9.0, 16.0),
            ("3.4", "Fase 3: Integración & APIs", "Bus de Eventos Kafka & Sincronización Real-Time", "SÍ", 7.0, 12.0, 22.0),
            ("3.5", "Fase 3: Integración & APIs", "Pipeline ETL de Reportería & Business Intelligence", "NO", 6.0, 10.0, 18.0),

            ("4.1", "Fase 4: Testing & Ciberseguridad", "Pruebas de Integración Continua Automatizadas", "NO", 5.0, 8.0, 15.0),
            ("4.2", "Fase 4: Testing & Ciberseguridad", "Pruebas de Carga, Rendimiento & Estrés (10k TPS)", "SÍ", 6.0, 10.0, 20.0),
            ("4.3", "Fase 4: Testing & Ciberseguridad", "Auditoría de Ciberseguridad & Pentesting Externo", "SÍ", 5.0, 9.0, 18.0),
            ("4.4", "Fase 4: Testing & Ciberseguridad", "Pruebas de Aceptación de Usuario (UAT Negocio)", "SÍ", 7.0, 12.0, 24.0),
            ("4.5", "Fase 4: Testing & Ciberseguridad", "Certificación Regulatoria de Seguridad & PCI-DSS", "SÍ", 6.0, 10.0, 22.0),

            ("5.1", "Fase 5: Migración & Go-Live", "Plan de Contingencia & Rollback Operativo", "NO", 3.0, 5.0, 10.0),
            ("5.2", "Fase 5: Migración & Go-Live", "Migración de Datos Históricos & Sanitización", "SÍ", 6.0, 11.0, 22.0),
            ("5.3", "Fase 5: Migración & Go-Live", "Simulación Mock Cutover en Entorno Staging", "SÍ", 4.0, 7.0, 15.0),
            ("5.4", "Fase 5: Migración & Go-Live", "Despliegue Go-Live en Producción & Switch DNS", "SÍ", 2.0, 4.0, 10.0),
            ("5.5", "Fase 5: Migración & Go-Live", "Soporte Hipercare 24/7 Post-Lanzamiento", "NO", 5.0, 10.0, 18.0),
        ]
    else:
        tasks = [
            ("1.1", "Phase 1: Architecture & Design", "Cloud Microservices Architecture Definition", "YES", 6.0, 10.0, 20.0),
            ("1.2", "Phase 1: Architecture & Design", "Data Modeling & IAM Security Governance", "NO", 4.0, 7.0, 14.0),
            ("1.3", "Phase 1: Architecture & Design", "API Specifications & OpenAPI Contracts", "YES", 5.0, 8.0, 17.0),
            ("1.4", "Phase 1: Architecture & Design", "Architecture Board Formal Approval", "YES", 3.0, 5.0, 12.0),
            ("1.5", "Phase 1: Architecture & Design", "IaC Terraform Infrastructure Provisioning", "NO", 4.0, 6.0, 12.0),

            ("2.1", "Phase 2: Core Development", "Core Transactional Reconciliation Engine", "YES", 12.0, 18.0, 32.0),
            ("2.2", "Phase 2: Core Development", "OAuth2 / OIDC Authentication Implementation", "NO", 5.0, 8.0, 15.0),
            ("2.3", "Phase 2: Core Development", "Payment Processing Module & Payment Gateways", "YES", 10.0, 16.0, 28.0),
            ("2.4", "Phase 2: Core Development", "Business Rules Engine & Order Routing Logic", "YES", 8.0, 14.0, 24.0),
            ("2.5", "Phase 2: Core Development", "Audit Logging & Distributed Tracing System", "NO", 4.0, 7.0, 13.0),

            ("3.1", "Phase 3: Integration & APIs", "Bidirectional ERP Connector (SAP/Oracle Core)", "YES", 10.0, 17.0, 30.0),
            ("3.2", "Phase 3: Integration & APIs", "Acquiring Banks & Gateway API Integration", "YES", 8.0, 14.0, 26.0),
            ("3.3", "Phase 3: Integration & APIs", "Salesforce CRM & Customer Care Connector", "NO", 5.0, 9.0, 16.0),
            ("3.4", "Phase 3: Integration & APIs", "Kafka Event Streaming & Real-Time Sync Bus", "YES", 7.0, 12.0, 22.0),
            ("3.5", "Phase 3: Integration & APIs", "BI Reporting Pipeline & ETL Data Warehouse", "NO", 6.0, 10.0, 18.0),

            ("4.1", "Phase 4: Testing & Cyber Security", "Automated Continuous Integration Testing Suite", "NO", 5.0, 8.0, 15.0),
            ("4.2", "Phase 4: Testing & Cyber Security", "Load, Performance & Stress Testing (10k TPS)", "YES", 6.0, 10.0, 20.0),
            ("4.3", "Phase 4: Testing & Cyber Security", "Cybersecurity Audit & External Pentesting", "YES", 5.0, 9.0, 18.0),
            ("4.4", "Phase 4: Testing & Cyber Security", "User Acceptance Testing (UAT Sign-off)", "YES", 7.0, 12.0, 24.0),
            ("4.5", "Phase 4: Testing & Cyber Security", "PCI-DSS Compliance Certification Sign-off", "YES", 6.0, 10.0, 22.0),

            ("5.1", "Phase 5: Migration & Cutover", "Rollback Protocol & Contingency Runbook", "NO", 3.0, 5.0, 10.0),
            ("5.2", "Phase 5: Migration & Cutover", "Legacy Data Sanitization & Cutover Migration", "YES", 6.0, 11.0, 22.0),
            ("5.3", "Phase 5: Migration & Cutover", "Mock Cutover Rehearsal on Staging Cluster", "YES", 4.0, 7.0, 15.0),
            ("5.4", "Phase 5: Migration & Cutover", "Production Go-Live Cutover & DNS Switch", "YES", 2.0, 4.0, 10.0),
            ("5.5", "Phase 5: Migration & Cutover", "24/7 Hypercare Post-Launch Operations", "NO", 5.0, 10.0, 18.0),
        ]
    return tasks


def create_pert_workbook(lang='ES'):
    """Construye el libro completo con 4 pestañas funcionales."""
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    default_sheet = wb.active

    # Nombres de hojas según idioma
    if lang == 'ES':
        t_dash = "Dashboard Ejecutivo"
        t_wbs = "Estimación WBS 3 Puntos"
        t_crit = "Camino Crítico & Análisis Z"
        t_buf = "Buffers & Plan de Aceleración"
    else:
        t_dash = "Executive Dashboard"
        t_wbs = "WBS 3-Point Estimation"
        t_crit = "Critical Path & Z-Score"
        t_buf = "Schedule Buffers & Crashing"

    ws_dash = wb.create_sheet(title=t_dash)
    ws_wbs = wb.create_sheet(title=t_wbs)
    ws_crit = wb.create_sheet(title=t_crit)
    ws_buf = wb.create_sheet(title=t_buf)
    wb.remove(default_sheet)

    # -------------------------------------------------------------------------
    # TAB 2: ESTIMACIÓN WBS 3 PUNTOS
    # -------------------------------------------------------------------------
    # Cabecera institucional
    ws_wbs.merge_cells("A1:M1")
    ws_wbs["A1"] = "DATALARIA | EXECUTIVE DECISION PACK · PMBOK 7th / STOCHASTIC MANAGEMENT" if lang == 'ES' else "DATALARIA | EXECUTIVE DECISION PACK · PMBOK 7th / STOCHASTIC MANAGEMENT"
    ws_wbs["A1"].font = font_subtitle
    ws_wbs["A1"].fill = fill_dark
    ws_wbs["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_wbs.merge_cells("A2:M2")
    ws_wbs["A2"] = "MODELO DE ESTIMACIÓN ESTOCÁSTICA DE 3 PUNTOS (DISTRIBUCIÓN BETA)" if lang == 'ES' else "STOCHASTIC 3-POINT ESTIMATION MODEL (BETA DISTRIBUTION)"
    ws_wbs["A2"].font = font_title
    ws_wbs["A2"].fill = fill_dark
    ws_wbs["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_wbs.merge_cells("A3:M3")
    ws_wbs["A3"] = "Cálculo riguroso de Duración Esperada μ, Desviación Estándar σ y Varianza σ² para cada paquete de trabajo según la WBS." if lang == 'ES' else "Rigorous calculation of Expected Duration μ, Standard Deviation σ, and Variance σ² for each WBS work package."
    ws_wbs["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="94A3B8")
    ws_wbs["A3"].fill = fill_dark
    ws_wbs["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Fila de espacio en fila 4
    ws_wbs.row_dimensions[1].height = 20
    ws_wbs.row_dimensions[2].height = 28
    ws_wbs.row_dimensions[3].height = 20
    ws_wbs.row_dimensions[4].height = 10

    # Guía de colores para el usuario
    ws_wbs.merge_cells("A5:M5")
    guide_text = "🔵 Celdas Blancas: Entradas editables del usuario (Nombre, Camino Crítico, Optimista, Más Probable, Pesimista). | 🔘 Celdas Grises: Cálculos protegidos." if lang == 'ES' else "🔵 White Cells: User editable inputs (Name, Critical Path, Optimistic, Most Likely, Pessimistic). | 🔘 Gray Cells: Protected formulas."
    ws_wbs["A5"] = guide_text
    ws_wbs["A5"].font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_BLUE_ACCENT)
    ws_wbs["A5"].fill = fill_blue_card
    ws_wbs["A5"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Cabeceras de tabla WBS (Fila 7)
    wbs_headers_es = [
        ("A7", "Código WBS"),
        ("B7", "Fase del Proyecto"),
        ("C7", "Paquete de Trabajo / Tarea"),
        ("D7", "Camino Crítico"),
        ("E7", "Optimista (o) [días]"),
        ("F7", "Más Probable (m) [días]"),
        ("G7", "Pesimista (p) [días]"),
        ("H7", "Media PERT Beta (μ)"),
        ("I7", "Media Triangular"),
        ("J7", "Desv. Estándar (σ)"),
        ("K7", "Varianza (σ²)"),
        ("L7", "Rango (p - o)"),
        ("M7", "Sesgo / Asimetría")
    ]
    wbs_headers_en = [
        ("A7", "WBS Code"),
        ("B7", "Project Phase"),
        ("C7", "Work Package / Deliverable"),
        ("D7", "Critical Path"),
        ("E7", "Optimistic (o) [days]"),
        ("F7", "Most Likely (m) [days]"),
        ("G7", "Pessimistic (p) [days]"),
        ("H7", "Expected PERT (μ)"),
        ("I7", "Triangular Mean"),
        ("J7", "Std Deviation (σ)"),
        ("K7", "Variance (σ²)"),
        ("L7", "Range (p - o)"),
        ("M7", "Skewness Ratio")
    ]
    headers_wbs = wbs_headers_es if lang == 'ES' else wbs_headers_en

    ws_wbs.row_dimensions[7].height = 26
    for cell_id, text in headers_wbs:
        cell = ws_wbs[cell_id]
        cell.value = text
        cell.font = font_tbl_header
        cell.fill = fill_med
        cell.alignment = align_center
        cell.border = border_cell

    # Insertar datos de tareas (Filas 8 a 32)
    tasks_data = build_wbs_data(lang)
    input_ranges_wbs = []

    for idx, t in enumerate(tasks_data, start=8):
        ws_wbs.row_dimensions[idx].height = 20
        code, phase, name, crit, o_val, m_val, p_val = t

        ws_wbs[f"A{idx}"] = code
        ws_wbs[f"A{idx}"].font = font_body_bold
        ws_wbs[f"A{idx}"].alignment = align_center
        ws_wbs[f"A{idx}"].fill = fill_formula
        ws_wbs[f"A{idx}"].border = border_cell

        ws_wbs[f"B{idx}"] = phase
        ws_wbs[f"B{idx}"].font = font_body
        ws_wbs[f"B{idx}"].alignment = align_left
        ws_wbs[f"B{idx}"].fill = fill_formula
        ws_wbs[f"B{idx}"].border = border_cell

        ws_wbs[f"C{idx}"] = name
        ws_wbs[f"C{idx}"].font = font_body
        ws_wbs[f"C{idx}"].alignment = align_left
        ws_wbs[f"C{idx}"].fill = fill_input
        ws_wbs[f"C{idx}"].border = border_cell
        input_ranges_wbs.append(f"C{idx}")

        ws_wbs[f"D{idx}"] = crit
        ws_wbs[f"D{idx}"].font = font_body_bold
        ws_wbs[f"D{idx}"].alignment = align_center
        ws_wbs[f"D{idx}"].fill = fill_input
        ws_wbs[f"D{idx}"].border = border_cell
        input_ranges_wbs.append(f"D{idx}")

        ws_wbs[f"E{idx}"] = o_val
        ws_wbs[f"E{idx}"].font = font_input
        ws_wbs[f"E{idx}"].alignment = align_center
        ws_wbs[f"E{idx}"].number_format = "0.0"
        ws_wbs[f"E{idx}"].fill = fill_input
        ws_wbs[f"E{idx}"].border = border_cell
        input_ranges_wbs.append(f"E{idx}")

        ws_wbs[f"F{idx}"] = m_val
        ws_wbs[f"F{idx}"].font = font_input
        ws_wbs[f"F{idx}"].alignment = align_center
        ws_wbs[f"F{idx}"].number_format = "0.0"
        ws_wbs[f"F{idx}"].fill = fill_input
        ws_wbs[f"F{idx}"].border = border_cell
        input_ranges_wbs.append(f"F{idx}")

        ws_wbs[f"G{idx}"] = p_val
        ws_wbs[f"G{idx}"].font = font_input
        ws_wbs[f"G{idx}"].alignment = align_center
        ws_wbs[f"G{idx}"].number_format = "0.0"
        ws_wbs[f"G{idx}"].fill = fill_input
        ws_wbs[f"G{idx}"].border = border_cell
        input_ranges_wbs.append(f"G{idx}")

        # Fórmulas estadísticas estándar
        # μ = (o + 4m + p) / 6
        ws_wbs[f"H{idx}"] = f"=(E{idx} + 4*F{idx} + G{idx}) / 6"
        ws_wbs[f"H{idx}"].font = font_body_bold
        ws_wbs[f"H{idx}"].alignment = align_center
        ws_wbs[f"H{idx}"].number_format = "0.0"
        ws_wbs[f"H{idx}"].fill = fill_formula
        ws_wbs[f"H{idx}"].border = border_cell

        # μ_tri = (o + m + p) / 3
        ws_wbs[f"I{idx}"] = f"=(E{idx} + F{idx} + G{idx}) / 3"
        ws_wbs[f"I{idx}"].font = font_body
        ws_wbs[f"I{idx}"].alignment = align_center
        ws_wbs[f"I{idx}"].number_format = "0.0"
        ws_wbs[f"I{idx}"].fill = fill_formula
        ws_wbs[f"I{idx}"].border = border_cell

        # σ = (p - o) / 6
        ws_wbs[f"J{idx}"] = f"=(G{idx} - E{idx}) / 6"
        ws_wbs[f"J{idx}"].font = font_body
        ws_wbs[f"J{idx}"].alignment = align_center
        ws_wbs[f"J{idx}"].number_format = "0.00"
        ws_wbs[f"J{idx}"].fill = fill_formula
        ws_wbs[f"J{idx}"].border = border_cell

        # σ² = ((p - o) / 6)²
        ws_wbs[f"K{idx}"] = f"=(J{idx})^2"
        ws_wbs[f"K{idx}"].font = font_body
        ws_wbs[f"K{idx}"].alignment = align_center
        ws_wbs[f"K{idx}"].number_format = "0.00"
        ws_wbs[f"K{idx}"].fill = fill_formula
        ws_wbs[f"K{idx}"].border = border_cell

        # Rango = p - o
        ws_wbs[f"L{idx}"] = f"=G{idx} - E{idx}"
        ws_wbs[f"L{idx}"].font = font_body
        ws_wbs[f"L{idx}"].alignment = align_center
        ws_wbs[f"L{idx}"].number_format = "0.0"
        ws_wbs[f"L{idx}"].fill = fill_formula
        ws_wbs[f"L{idx}"].border = border_cell

        # Sesgo / Asimetría
        ws_wbs[f"M{idx}"] = f'=IF((G{idx}-E{idx})>0, ROUND((G{idx}+E{idx}-2*F{idx})/(G{idx}-E{idx}), 2), 0)'
        ws_wbs[f"M{idx}"].font = font_body
        ws_wbs[f"M{idx}"].alignment = align_center
        ws_wbs[f"M{idx}"].number_format = "0.00"
        ws_wbs[f"M{idx}"].fill = fill_formula
        ws_wbs[f"M{idx}"].border = border_cell

    # Filas de resumen agregado (Filas 34 a 37)
    crit_tag = "SÍ" if lang == 'ES' else "YES"

    # Filas de resumen agregado (Filas 34 a 38)
    crit_tag = "SÍ" if lang == 'ES' else "YES"

    ws_wbs.row_dimensions[34].height = 24
    ws_wbs.row_dimensions[35].height = 24
    ws_wbs.row_dimensions[36].height = 26
    ws_wbs.row_dimensions[37].height = 26
    ws_wbs.row_dimensions[38].height = 26

    # Fila 34: Total Tareas
    ws_wbs["B34"] = "TOTAL PAQUETES DE TRABAJO WBS" if lang == 'ES' else "TOTAL WBS WORK PACKAGES"
    ws_wbs["B34"].font = font_body_bold
    ws_wbs["B34"].border = border_cell
    ws_wbs["C34"] = "=COUNTA(A8:A32)"
    ws_wbs["C34"].font = font_body_bold
    ws_wbs["C34"].alignment = align_center
    ws_wbs["C34"].border = border_cell

    # Fila 35: Tareas en Camino Crítico
    ws_wbs["B35"] = "TOTAL TAREAS EN CAMINO CRÍTICO" if lang == 'ES' else "TOTAL CRITICAL PATH TASKS"
    ws_wbs["B35"].font = font_body_bold
    ws_wbs["B35"].border = border_cell
    ws_wbs["C35"] = f'=COUNTIF(D8:D32, "{crit_tag}")'
    ws_wbs["C35"].font = font_body_bold
    ws_wbs["C35"].alignment = align_center
    ws_wbs["C35"].border = border_cell

    # Fila 36: Suma de Duraciones del Camino Crítico (μ_proj)
    ws_wbs.merge_cells("E36:G36")
    ws_wbs["E36"] = "DURACIÓN ESPERADA CRÍTICA (Σ μ_crit):" if lang == 'ES' else "CRITICAL EXPECTED DURATION (Σ μ_crit):"
    ws_wbs["E36"].font = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_BLUE_ACCENT)
    ws_wbs["E36"].alignment = Alignment(horizontal="right", vertical="center")
    for col_c in ["E", "F", "G"]:
        ws_wbs[f"{col_c}36"].border = border_cell
    ws_wbs["H36"] = f'=SUMIF(D8:D32, "{crit_tag}", H8:H32)'
    ws_wbs["H36"].font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_BLUE_ACCENT)
    ws_wbs["H36"].number_format = "0.0"
    ws_wbs["H36"].alignment = align_center
    ws_wbs["H36"].fill = fill_blue_card
    ws_wbs["H36"].border = border_card

    # Fila 37: Suma de Varianzas del Camino Crítico (Σ σ²_crit)
    ws_wbs.merge_cells("E37:G37")
    ws_wbs["E37"] = "VARIANZA COMBINADA (Σ σ²_crit):" if lang == 'ES' else "COMBINED VARIANCE (Σ σ²_crit):"
    ws_wbs["E37"].font = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_NAVY_DARK)
    ws_wbs["E37"].alignment = Alignment(horizontal="right", vertical="center")
    for col_c in ["E", "F", "G"]:
        ws_wbs[f"{col_c}37"].border = border_cell
    ws_wbs["H37"] = f'=SUMIF(D8:D32, "{crit_tag}", K8:K32)'
    ws_wbs["H37"].font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY_DARK)
    ws_wbs["H37"].number_format = "0.00"
    ws_wbs["H37"].alignment = align_center
    ws_wbs["H37"].fill = fill_formula
    ws_wbs["H37"].border = border_cell
    ws_wbs["K37"] = "=H37"

    # Fila 38: Desviación global (σ_proj)
    ws_wbs.merge_cells("E38:G38")
    ws_wbs["E38"] = "DESVIACIÓN GLOBAL (σ_proj):" if lang == 'ES' else "GLOBAL STD DEV (σ_proj):"
    ws_wbs["E38"].font = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_NAVY_DARK)
    ws_wbs["E38"].alignment = Alignment(horizontal="right", vertical="center")
    for col_c in ["E", "F", "G"]:
        ws_wbs[f"{col_c}38"].border = border_cell
    ws_wbs["H38"] = "=SQRT(H37)"
    ws_wbs["H38"].font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_NAVY_DARK)
    ws_wbs["H38"].number_format = "0.00"
    ws_wbs["H38"].alignment = align_center
    ws_wbs["H38"].fill = fill_blue_card
    ws_wbs["H38"].border = border_card

    # Ajuste de anchos de columna en WBS
    col_widths_wbs = {
        "A": 12, "B": 28, "C": 44, "D": 16, "E": 18, "F": 20,
        "G": 18, "H": 20, "I": 18, "J": 18, "K": 16, "L": 16, "M": 16
    }
    for col, width in col_widths_wbs.items():
        ws_wbs.column_dimensions[col].width = width

    # -------------------------------------------------------------------------
    # TAB 3: CAMINO CRÍTICO & ANÁLISIS Z (CENTRAL LIMIT THEOREM)
    # -------------------------------------------------------------------------
    ws_crit.merge_cells("A1:G1")
    ws_crit["A1"] = "DATALARIA | EXECUTIVE DECISION PACK · PMBOK 7th / STOCHASTIC MANAGEMENT"
    ws_crit["A1"].font = font_subtitle
    ws_crit["A1"].fill = fill_dark
    ws_crit["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_crit.merge_cells("A2:G2")
    ws_crit["A2"] = "AGREGACIÓN ESTOCÁSTICA DEL CAMINO CRÍTICO & ANÁLISIS DE CERTIDUMBRE Z" if lang == 'ES' else "CRITICAL PATH STOCHASTIC AGGREGATION & Z-SCORE CERTAINTY ANALYSIS"
    ws_crit["A2"].font = font_title
    ws_crit["A2"].fill = fill_dark
    ws_crit["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_crit.merge_cells("A3:G3")
    ws_crit["A3"] = "Teorema del Límite Central (CLT): Agregación cuadrática de varianzas e inversión de probabilidad normal para el Consejo de Administración." if lang == 'ES' else "Central Limit Theorem (CLT): Quadratic variance aggregation and normal probability inversion for Board Decision-Making."
    ws_crit["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="94A3B8")
    ws_crit["A3"].fill = fill_dark
    ws_crit["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_crit.row_dimensions[1].height = 20
    ws_crit.row_dimensions[2].height = 28
    ws_crit.row_dimensions[3].height = 20
    ws_crit.row_dimensions[4].height = 10

    # Sección 1: Parámetros Globales del Proyecto (Filas 6 a 12)
    ws_crit["A6"] = "1. PARÁMETROS AGREGADOS DEL CRONOGRAMA GLOBAL (CLT)" if lang == 'ES' else "1. GLOBAL SCHEDULE AGGREGATED PARAMETERS (CLT)"
    ws_crit["A6"].font = font_section

    crit_summary_labels_es = [
        ("A7", "Duración Esperada del Proyecto (μ_total):", f"='{t_wbs}'!H36", "días hábiles", "0.0"),
        ("A8", "Varianza Combinada del Proyecto (σ²_total):", f"='{t_wbs}'!H37", "días²", "0.00"),
        ("A9", "Desviación Estándar Global (σ_total):", f"=SQRT(C8)", "días hábiles", "0.00"),
        ("A10", "Intervalo de Confianza 68.3% (μ ± 1σ):", f"=C7-C9", f"=C7+C9", "0.0"),
        ("A11", "Intervalo de Confianza 90.0% (μ ± 1.645σ):", f"=C7-1.645*C9", f"=C7+1.645*C9", "0.0"),
        ("A12", "Intervalo de Confianza 95.4% (μ ± 2σ):", f"=C7-2*C9", f"=C7+2*C9", "0.0"),
    ]
    crit_summary_labels_en = [
        ("A7", "Expected Project Duration (μ_total):", f"='{t_wbs}'!H36", "business days", "0.0"),
        ("A8", "Combined Project Variance (σ²_total):", f"='{t_wbs}'!H37", "days²", "0.00"),
        ("A9", "Global Project Std Deviation (σ_total):", f"=SQRT(C8)", "business days", "0.00"),
        ("A10", "68.3% Confidence Interval (μ ± 1σ):", f"=C7-C9", f"=C7+C9", "0.0"),
        ("A11", "90.0% Confidence Interval (μ ± 1.645σ):", f"=C7-1.645*C9", f"=C7+1.645*C9", "0.0"),
        ("A12", "95.4% Confidence Interval (μ ± 2σ):", f"=C7-2*C9", f"=C7+2*C9", "0.0"),
    ]
    summary_labels = crit_summary_labels_es if lang == 'ES' else crit_summary_labels_en

    for item in summary_labels:
        cell_lbl = item[0]
        row_num = int(cell_lbl[1:])
        ws_crit.row_dimensions[row_num].height = 24

        ws_crit[cell_lbl] = item[1]
        ws_crit[cell_lbl].font = font_body_bold
        ws_crit[cell_lbl].alignment = align_left

        if len(item) == 5 and row_num <= 9:
            ws_crit[f"C{row_num}"] = item[2]
            ws_crit[f"C{row_num}"].font = font_body_bold
            ws_crit[f"C{row_num}"].fill = fill_blue_card
            ws_crit[f"C{row_num}"].alignment = align_center
            ws_crit[f"C{row_num}"].number_format = item[4]
            ws_crit[f"C{row_num}"].border = border_card

            ws_crit[f"D{row_num}"] = item[3]
            ws_crit[f"D{row_num}"].font = Font(name=FONT_NAME, size=9, italic=True, color="64748B")
            ws_crit[f"D{row_num}"].alignment = align_left
        else:
            ws_crit[f"C{row_num}"] = item[2]
            ws_crit[f"C{row_num}"].font = font_body
            ws_crit[f"C{row_num}"].alignment = align_center
            ws_crit[f"C{row_num}"].number_format = "0.0"
            ws_crit[f"C{row_num}"].fill = fill_formula
            ws_crit[f"C{row_num}"].border = border_cell

            ws_crit[f"D{row_num}"] = " a " if lang == 'ES' else " to "
            ws_crit[f"D{row_num}"].font = font_body
            ws_crit[f"D{row_num}"].alignment = align_center

            ws_crit[f"E{row_num}"] = item[3]
            ws_crit[f"E{row_num}"].font = font_body
            ws_crit[f"E{row_num}"].alignment = align_center
            ws_crit[f"E{row_num}"].number_format = "0.0"
            ws_crit[f"E{row_num}"].fill = fill_formula
            ws_crit[f"E{row_num}"].border = border_cell

    # Sección 2: Tabla de Escenarios de Confianza (Filas 15 a 26)
    ws_crit["A15"] = "2. TABLA DE ESCENARIOS DE CONFIANZA & COMPROMISOS DIRECTIVOS" if lang == 'ES' else "2. CONFIDENCE SCENARIOS & EXECUTIVE COMMITMENT TABLE"
    ws_crit["A15"].font = font_section

    conf_headers_es = [
        ("A16", "Nivel Confianza (P)"),
        ("B16", "Z-Score Estándar"),
        ("C16", "Duración Proyecto [días]"),
        ("D16", "Colchón / Buffer Requerido"),
        ("E16", "% Buffer / Base μ"),
        ("F16", "Perfil de Apetito de Riesgo"),
        ("G16", "Recomendación de Gobierno Corporativo")
    ]
    conf_headers_en = [
        ("A16", "Confidence Level (P)"),
        ("B16", "Standard Z-Score"),
        ("C16", "Required Duration [days]"),
        ("D16", "Required Schedule Buffer"),
        ("E16", "% Buffer / Base μ"),
        ("F16", "Risk Appetite Profile"),
        ("G16", "Corporate Governance Recommendation")
    ]
    headers_conf = conf_headers_es if lang == 'ES' else conf_headers_en

    ws_crit.row_dimensions[16].height = 28
    for cid, txt in headers_conf:
        c = ws_crit[cid]
        c.value = txt
        c.font = font_tbl_header
        c.fill = fill_med
        c.alignment = align_center
        c.border = border_cell

    # Datos de percentiles: 50%, 60%, 70%, 75%, 80%, 85%, 90%, 95%, 99%
    percentiles_es = [
        (0.50, "P50 (Mediana Base)", "Moneda al aire: 50% probabilidad de retraso.", "INSUFICIENTE para compromisos directivos o contractuales."),
        (0.60, "P60 (Riesgo Elevado)", "Muy agresivo. Fuerte exposición a penalizaciones.", "Solo viable en prototipos internos sin impacto en cliente."),
        (0.70, "P70 (Agresivo)", "Riesgo moderado-alto. Requiere equipo sénior.", "Válido si existe holgura financiera para absorber retraso."),
        (0.75, "P75 (Aceptable)", "Umbral operativo mínimo para releases ágiles.", "Exige revisión semanal del camino crítico."),
        (0.80, "P80 (Prudente PMBOK)", "Estándar recomendado por PMBOK para proyectos comerciales.", "Alineación equilibrada entre coste de colchón y certeza."),
        (0.85, "P85 (Alta Confianza)", "Compromiso sólido con bajo margen de desvío.", "Recomendado para contratos con penalización moderada."),
        (0.90, "P90 (Estándar Consejo C-Level)", "⭐⭐ UMBRAL OFICIAL CONSEJO DE ADMINISTRACIÓN", "Certidumbre exigible ante el Board y auditoría externa."),
        (0.95, "P95 (Misión Crítica)", "Alta certidumbre para proyectos de infraestructura core.", "Obligatorio en entornos bancarios, cloud y salud."),
        (0.99, "P99 (Tolerancia Cero)", "Certeza casi absoluta. Máxima dotación de colchón.", "Aeroespacial, defensa y sistemas vitales donde el fallo es fatal.")
    ]
    percentiles_en = [
        (0.50, "P50 (Base Median)", "Coin toss: 50% probability of project delay.", "INSUFICIENT for executive or contractual commitments."),
        (0.60, "P60 (High Risk)", "Highly aggressive. Strong exposure to penalties.", "Only viable for internal prototypes with no client impact."),
        (0.70, "P70 (Aggressive)", "Moderate-high risk. Requires senior delivery team.", "Acceptable only if budget exists to absorb variance."),
        (0.75, "P75 (Acceptable)", "Minimum operational threshold for agile releases.", "Requires weekly critical path burn-rate reviews."),
        (0.80, "P80 (Prudent PMBOK)", "Standard PMBOK recommended baseline.", "Balanced alignment between buffer cost and reliability."),
        (0.85, "P85 (High Confidence)", "Solid delivery commitment with low disruption risk.", "Recommended for contracts with moderate penalties."),
        (0.90, "P90 (Board C-Level Standard)", "⭐⭐ OFFICIAL BOARD OF DIRECTORS THRESHOLD", "Expected certainty for Board sign-off and external audit."),
        (0.95, "P95 (Mission Critical)", "High certainty for core financial/cloud infrastructure.", "Mandatory in banking, regulated SaaS and healthcare."),
        (0.99, "P99 (Zero Tolerance)", "Near absolute certainty. Maximum contingency sizing.", "Aerospace, defense and life-critical systems.")
    ]
    pct_data = percentiles_es if lang == 'ES' else percentiles_en

    for i, p in enumerate(pct_data, start=17):
        ws_crit.row_dimensions[i].height = 36
        prob, label, risk_prof, gov_rec = p

        # Col A: Probabilidad
        ws_crit[f"A{i}"] = prob
        ws_crit[f"A{i}"].number_format = "0.0%"
        ws_crit[f"A{i}"].font = font_body_bold
        ws_crit[f"A{i}"].alignment = align_center
        ws_crit[f"A{i}"].border = border_cell

        # Col B: Z-Score
        ws_crit[f"B{i}"] = f"=NORMSINV(A{i})"
        ws_crit[f"B{i}"].number_format = "0.00"
        ws_crit[f"B{i}"].font = font_body
        ws_crit[f"B{i}"].alignment = align_center
        ws_crit[f"B{i}"].border = border_cell

        # Col C: Duración Requerida
        ws_crit[f"C{i}"] = f"=NORMINV(A{i}, $C$7, $C$9)"
        ws_crit[f"C{i}"].number_format = "0.0"
        ws_crit[f"C{i}"].font = font_body_bold
        ws_crit[f"C{i}"].alignment = align_center
        ws_crit[f"C{i}"].border = border_cell

        # Col D: Buffer Requerido
        ws_crit[f"D{i}"] = f"=C{i} - $C$7"
        ws_crit[f"D{i}"].number_format = "+0.0;-0.0;0.0"
        ws_crit[f"D{i}"].font = font_body_bold
        ws_crit[f"D{i}"].alignment = align_center
        ws_crit[f"D{i}"].border = border_cell

        # Col E: % Buffer
        ws_crit[f"E{i}"] = f"=D{i} / $C$7"
        ws_crit[f"E{i}"].number_format = "0.0%"
        ws_crit[f"E{i}"].font = font_body
        ws_crit[f"E{i}"].alignment = align_center
        ws_crit[f"E{i}"].border = border_cell

        # Col F: Perfil
        ws_crit[f"F{i}"] = risk_prof
        ws_crit[f"F{i}"].font = font_body
        ws_crit[f"F{i}"].alignment = align_left
        ws_crit[f"F{i}"].border = border_cell

        # Col G: Recomendación
        ws_crit[f"G{i}"] = gov_rec
        ws_crit[f"G{i}"].font = font_body
        ws_crit[f"G{i}"].alignment = align_left
        ws_crit[f"G{i}"].border = border_cell

        # Resaltado especial de la fila P90 (i=23)
        if prob == 0.90:
            for col_c in ["A", "B", "C", "D", "E", "F", "G"]:
                ws_crit[f"{col_c}{i}"].fill = fill_green
                ws_crit[f"{col_c}{i}"].font = font_body_bold

    # Anchos de columna en Tab 3
    col_widths_crit = {
        "A": 48, "B": 18, "C": 24, "D": 24, "E": 18, "F": 48, "G": 58
    }
    for col, width in col_widths_crit.items():
        ws_crit.column_dimensions[col].width = width

    # -------------------------------------------------------------------------
    # TAB 4: BUFFERS & PLAN DE ACELERACIÓN (CRASHING & FAST-TRACKING)
    # -------------------------------------------------------------------------
    ws_buf.merge_cells("A1:K1")
    ws_buf["A1"] = "DATALARIA | EXECUTIVE DECISION PACK · PMBOK 7th / STOCHASTIC MANAGEMENT"
    ws_buf["A1"].font = font_subtitle
    ws_buf["A1"].fill = fill_dark
    ws_buf["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_buf.merge_cells("A2:K2")
    ws_buf["A2"] = "DIMENSIONAMIENTO DE BUFFERS & PLAN DE ACELERACIÓN (CRASHING / FAST-TRACKING)" if lang == 'ES' else "BUFFER SIZING & SCHEDULE ACCELERATION PLAN (CRASHING / FAST-TRACKING)"
    ws_buf["A2"].font = font_title
    ws_buf["A2"].fill = fill_dark
    ws_buf["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_buf.merge_cells("A3:K3")
    ws_buf["A3"] = "Optimización coste-tiempo para reducir la duración crítica mediante crashing presupuestario y colchones de proyecto según Teoría de Restricciones (TOC)." if lang == 'ES' else "Cost-time optimization to reduce critical path duration via budgetary crashing and project buffers according to Theory of Constraints (TOC)."
    ws_buf["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="94A3B8")
    ws_buf["A3"].fill = fill_dark
    ws_buf["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_buf.row_dimensions[1].height = 20
    ws_buf.row_dimensions[2].height = 28
    ws_buf.row_dimensions[3].height = 20
    ws_buf.row_dimensions[4].height = 10

    # Sección 1: Comparativa de Métodos de Dimensionamiento de Colchón (Filas 6 a 11)
    ws_buf["A6"] = "1. DIMENSIONAMIENTO CIENTÍFICO DE COLCHONES DE CONTINGENCIA (PROJECT BUFFERS)" if lang == 'ES' else "1. SCIENTIFIC PROJECT BUFFER SIZING METHODS"
    ws_buf["A6"].font = font_section

    buf_methods_es = [
        ("A7", "Método Tradicional Goldratt (50% Incertidumbre Pesimista):", f'=0.5 * (SUMIF(\'{t_wbs}\'!$D$8:$D$32, "SÍ", \'{t_wbs}\'!$G$8:$G$32) - SUMIF(\'{t_wbs}\'!$D$8:$D$32, "SÍ", \'{t_wbs}\'!$F$8:$F$32))', "días", "Cálculo heurístico lineal (sobreestima buffer)."),
        ("A8", "Método Raíz Cuadrada de Varianzas (RSS / CLT 2σ):", f"=2 * '{t_crit}'!$C$9", "días", "Basado en CLT: varianzas independientes agregadas."),
        ("A9", "Colchón Recomendado C-Level P90 (90% Certidumbre):", f"='{t_crit}'!$D$23", "días", "⭐ Estándar óptimo para aprobación de presupuesto."),
        ("A10", "Colchón de Alta Seguridad P95 (95% Certidumbre):", f"='{t_crit}'!$D$24", "días", "Recomendado para contratos con cláusulas penales severas."),
    ]
    buf_methods_en = [
        ("A7", "Goldratt Traditional Cut (50% Pessimistic Uncertainty):", f'=0.5 * (SUMIF(\'{t_wbs}\'!$D$8:$D$32, "YES", \'{t_wbs}\'!$G$8:$G$32) - SUMIF(\'{t_wbs}\'!$D$8:$D$32, "YES", \'{t_wbs}\'!$F$8:$F$32))', "days", "Linear heuristic formula (tends to overestimate buffer)."),
        ("A8", "Root-Sum-Square Method (RSS / CLT 2σ):", f"=2 * '{t_crit}'!$C$9", "days", "CLT based: quadratic aggregation of independent variances."),
        ("A9", "Recommended Board P90 Buffer (90% Certainty):", f"='{t_crit}'!$D$23", "days", "⭐ Optimal standard for executive Board approval."),
        ("A10", "High Security P95 Buffer (95% Certainty):", f"='{t_crit}'!$D$24", "days", "Recommended for contracts with liquidated damages."),
    ]
    methods_list = buf_methods_es if lang == 'ES' else buf_methods_en

    for b_item in methods_list:
        row_n = int(b_item[0][1:])
        ws_buf.row_dimensions[row_n].height = 26
        ws_buf[b_item[0]] = b_item[1]
        ws_buf[b_item[0]].font = font_body_bold
        ws_buf[b_item[0]].alignment = align_left

        ws_buf[f"D{row_n}"] = b_item[2]
        ws_buf[f"D{row_n}"].font = font_body_bold
        ws_buf[f"D{row_n}"].alignment = align_center
        ws_buf[f"D{row_n}"].number_format = "0.0"
        ws_buf[f"D{row_n}"].fill = fill_blue_card if row_n == 9 else fill_formula
        ws_buf[f"D{row_n}"].border = border_card if row_n == 9 else border_cell

        ws_buf[f"E{row_n}"] = b_item[3]
        ws_buf[f"E{row_n}"].font = Font(name=FONT_NAME, size=9, italic=True, color="64748B")
        ws_buf[f"E{row_n}"].alignment = align_left

        ws_buf[f"F{row_n}"] = b_item[4]
        ws_buf[f"F{row_n}"].font = font_body
        ws_buf[f"F{row_n}"].alignment = align_left

    # Sección 2: Matriz de Opciones de Crashing (Filas 14 a 26)
    ws_buf["A14"] = "2. MATRIZ DE CRASHING EN TAREAS CRÍTICAS (COSTE MARGINAL POR DÍA REDUCIDO)" if lang == 'ES' else "2. CRITICAL PATH CRASHING MATRIX (MARGINAL COST PER DAY REDUCED)"
    ws_buf["A14"].font = font_section

    crashing_headers_es = [
        ("A15", "WBS"),
        ("B15", "Tarea Crítica"),
        ("C15", "Duración Normal [d]"),
        ("D15", "Duración Crash [d]"),
        ("E15", "Reducción Máx [d]"),
        ("F15", "Coste Normal [€]"),
        ("G15", "Coste Crash [€]"),
        ("H15", "Sobrecolchón [€]"),
        ("I15", "Coste / Día Reducido [€/d]"),
        ("J15", "Prioridad Aceleración"),
        ("K15", "Decisión Comité / Board")
    ]
    crashing_headers_en = [
        ("A15", "WBS"),
        ("B15", "Critical Task"),
        ("C15", "Normal Duration [d]"),
        ("D15", "Crash Duration [d]"),
        ("E15", "Max Reduction [d]"),
        ("F15", "Normal Cost [$]"),
        ("G15", "Crash Cost [$]"),
        ("H15", "Delta Cost [$]"),
        ("I15", "Cost / Day Reduced [$/d]"),
        ("J15", "Crashing Priority"),
        ("K15", "Board Decision")
    ]
    c_hdrs = crashing_headers_es if lang == 'ES' else crashing_headers_en

    ws_buf.row_dimensions[15].height = 28
    for cid, txt in c_hdrs:
        c = ws_buf[cid]
        c.value = txt
        c.font = font_tbl_header
        c.fill = fill_med
        c.alignment = align_center
        c.border = border_cell

    # 10 tareas del camino crítico acelerables
    crash_tasks_es = [
        ("1.1", "Arquitectura Cloud & Microservicios", 11.0, 7.0, 18000, 24000, "Aprobado (Fase 1)"),
        ("1.3", "Especificación de APIs OpenAPI", 9.0, 6.0, 12000, 16500, "Aprobado (Fase 1)"),
        ("2.1", "Motor Transaccional Conciliación", 19.3, 13.0, 45000, 62000, "Aprobado (Fase 2)"),
        ("2.3", "Módulo de Procesamiento Pagos", 17.0, 12.0, 38000, 51000, "En Reserva (Fase 2)"),
        ("2.4", "Motor de Reglas de Negocio", 14.7, 10.0, 28000, 37500, "En Reserva (Fase 2)"),
        ("3.1", "Conector Bidireccional ERP SAP", 18.0, 13.0, 42000, 59000, "En Reserva (Fase 3)"),
        ("3.2", "Integración Pasarelas Bancarias", 15.0, 10.0, 32000, 44500, "Aprobado (Fase 3)"),
        ("4.2", "Pruebas de Carga & Estrés", 11.0, 7.0, 22000, 29000, "Aprobado (Fase 4)"),
        ("4.4", "Pruebas de Aceptación UAT", 13.0, 8.0, 20000, 29500, "En Reserva (Fase 4)"),
        ("5.2", "Migración de Datos Cutover", 12.0, 8.0, 25000, 36000, "En Reserva (Fase 5)")
    ]
    crash_tasks_en = [
        ("1.1", "Cloud Microservices Architecture", 11.0, 7.0, 18000, 24000, "Approved (Phase 1)"),
        ("1.3", "OpenAPI API Specification", 9.0, 6.0, 12000, 16500, "Approved (Phase 1)"),
        ("2.1", "Core Reconciliation Engine", 19.3, 13.0, 45000, 62000, "Approved (Phase 2)"),
        ("2.3", "Payment Processing Module", 17.0, 12.0, 38000, 51000, "On Hold (Phase 2)"),
        ("2.4", "Business Rules Engine", 14.7, 10.0, 28000, 37500, "On Hold (Phase 2)"),
        ("3.1", "Bidirectional SAP Connector", 18.0, 13.0, 42000, 59000, "On Hold (Phase 3)"),
        ("3.2", "Acquiring Banking Gateways", 15.0, 10.0, 32000, 44500, "Approved (Phase 3)"),
        ("4.2", "Stress & Load Testing (10k TPS)", 11.0, 7.0, 22000, 29000, "Approved (Phase 4)"),
        ("4.4", "User Acceptance Testing (UAT)", 13.0, 8.0, 20000, 29500, "On Hold (Phase 4)"),
        ("5.2", "Data Migration & Cutover", 12.0, 8.0, 25000, 36000, "On Hold (Phase 5)")
    ]
    c_list = crash_tasks_es if lang == 'ES' else crash_tasks_en
    input_ranges_buf = []

    curr_format = "#,##0 €" if lang == 'ES' else "$#,##0"

    for idx, row_data in enumerate(c_list, start=16):
        ws_buf.row_dimensions[idx].height = 24
        wbs_c, name_c, dur_norm, dur_crsh, cst_norm, cst_crsh, dec_brd = row_data

        ws_buf[f"A{idx}"] = wbs_c
        ws_buf[f"A{idx}"].font = font_body_bold
        ws_buf[f"A{idx}"].alignment = align_center
        ws_buf[f"A{idx}"].border = border_cell

        ws_buf[f"B{idx}"] = name_c
        ws_buf[f"B{idx}"].font = font_body
        ws_buf[f"B{idx}"].alignment = align_left
        ws_buf[f"B{idx}"].border = border_cell

        ws_buf[f"C{idx}"] = dur_norm
        ws_buf[f"C{idx}"].font = font_body
        ws_buf[f"C{idx}"].alignment = align_center
        ws_buf[f"C{idx}"].number_format = "0.0"
        ws_buf[f"C{idx}"].border = border_cell

        ws_buf[f"D{idx}"] = dur_crsh
        ws_buf[f"D{idx}"].font = font_input
        ws_buf[f"D{idx}"].alignment = align_center
        ws_buf[f"D{idx}"].number_format = "0.0"
        ws_buf[f"D{idx}"].fill = fill_input
        ws_buf[f"D{idx}"].border = border_cell
        input_ranges_buf.append(f"D{idx}")

        ws_buf[f"E{idx}"] = f"=C{idx} - D{idx}"
        ws_buf[f"E{idx}"].font = font_body_bold
        ws_buf[f"E{idx}"].alignment = align_center
        ws_buf[f"E{idx}"].number_format = "0.0"
        ws_buf[f"E{idx}"].border = border_cell

        ws_buf[f"F{idx}"] = cst_norm
        ws_buf[f"F{idx}"].font = font_input
        ws_buf[f"F{idx}"].alignment = align_right
        ws_buf[f"F{idx}"].number_format = curr_format
        ws_buf[f"F{idx}"].fill = fill_input
        ws_buf[f"F{idx}"].border = border_cell
        input_ranges_buf.append(f"F{idx}")

        ws_buf[f"G{idx}"] = cst_crsh
        ws_buf[f"G{idx}"].font = font_input
        ws_buf[f"G{idx}"].alignment = align_right
        ws_buf[f"G{idx}"].number_format = curr_format
        ws_buf[f"G{idx}"].fill = fill_input
        ws_buf[f"G{idx}"].border = border_cell
        input_ranges_buf.append(f"G{idx}")

        ws_buf[f"H{idx}"] = f"=G{idx} - F{idx}"
        ws_buf[f"H{idx}"].font = font_body
        ws_buf[f"H{idx}"].alignment = align_right
        ws_buf[f"H{idx}"].number_format = curr_format
        ws_buf[f"H{idx}"].border = border_cell

        # Coste marginal por día reducido = (Crash Cost - Normal Cost) / Reducción Máxima
        ws_buf[f"I{idx}"] = f"=IF(E{idx}>0, H{idx}/E{idx}, 0)"
        ws_buf[f"I{idx}"].font = font_body_bold
        ws_buf[f"I{idx}"].alignment = align_right
        ws_buf[f"I{idx}"].number_format = curr_format
        ws_buf[f"I{idx}"].border = border_cell

        # Prioridad de Aceleración (RANK: menor coste marginal es prioridad 1)
        ws_buf[f"J{idx}"] = f"=RANK(I{idx}, $I$16:$I$25, 1)"
        ws_buf[f"J{idx}"].font = font_body_bold
        ws_buf[f"J{idx}"].alignment = align_center
        ws_buf[f"J{idx}"].border = border_cell

        ws_buf[f"K{idx}"] = dec_brd
        ws_buf[f"K{idx}"].font = font_input
        ws_buf[f"K{idx}"].alignment = align_center
        ws_buf[f"K{idx}"].fill = fill_input
        ws_buf[f"K{idx}"].border = border_cell
        input_ranges_buf.append(f"K{idx}")

    # Totales de Crashing en Fila 27
    ws_buf.row_dimensions[27].height = 26
    ws_buf["B27"] = "TOTAL POTENCIAL DE CRASHING DISPONIBLE:" if lang == 'ES' else "TOTAL AVAILABLE CRASHING CAPACITY:"
    ws_buf["B27"].font = font_body_bold
    ws_buf["B27"].alignment = align_right

    ws_buf["E27"] = "=SUM(E16:E25)"
    ws_buf["E27"].font = font_body_bold
    ws_buf["E27"].alignment = align_center
    ws_buf["E27"].number_format = "0.0"
    ws_buf["E27"].fill = fill_blue_card
    ws_buf["E27"].border = border_card

    ws_buf["F27"] = "=SUM(F16:F25)"
    ws_buf["F27"].font = font_body_bold
    ws_buf["F27"].alignment = align_right
    ws_buf["F27"].number_format = curr_format
    ws_buf["F27"].border = border_cell

    ws_buf["G27"] = "=SUM(G16:G25)"
    ws_buf["G27"].font = font_body_bold
    ws_buf["G27"].alignment = align_right
    ws_buf["G27"].number_format = curr_format
    ws_buf["G27"].border = border_cell

    ws_buf["H27"] = "=SUM(H16:H25)"
    ws_buf["H27"].font = font_body_bold
    ws_buf["H27"].alignment = align_right
    ws_buf["H27"].number_format = curr_format
    ws_buf["H27"].fill = fill_formula
    ws_buf["H27"].border = border_cell

    # Anchos de columna en Tab 4
    col_widths_buf = {
        "A": 48, "B": 38, "C": 20, "D": 18, "E": 18,
        "F": 20, "G": 20, "H": 20, "I": 24, "J": 20, "K": 26
    }
    for col, width in col_widths_buf.items():
        ws_buf.column_dimensions[col].width = width

    # -------------------------------------------------------------------------
    # TAB 1: DASHBOARD EJECUTIVO (EXECUTIVE DASHBOARD)
    # -------------------------------------------------------------------------
    ws_dash.merge_cells("A1:K1")
    ws_dash["A1"] = "DATALARIA | EXECUTIVE DECISION PACK · PMBOK 7th / STOCHASTIC MANAGEMENT"
    ws_dash["A1"].font = font_subtitle
    ws_dash["A1"].fill = fill_dark
    ws_dash["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_dash.merge_cells("A2:K2")
    ws_dash["A2"] = "DASHBOARD EJECUTIVO: ESTIMACIÓN PERT & RIESGO DE CRONOGRAMA" if lang == 'ES' else "EXECUTIVE DASHBOARD: PERT ESTIMATION & SCHEDULE RISK"
    ws_dash["A2"].font = font_title
    ws_dash["A2"].fill = fill_dark
    ws_dash["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_dash.merge_cells("A3:K3")
    ws_dash["A3"] = "Control estocástico del camino crítico, calculadora interactiva de fecha límite y análisis de certidumbre para el Comité de Dirección." if lang == 'ES' else "Stochastic critical path control, interactive deadline calculator, and Board-level certainty analytics."
    ws_dash["A3"].font = Font(name=FONT_NAME, size=9.5, italic=True, color="94A3B8")
    ws_dash["A3"].fill = fill_dark
    ws_dash["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_dash.row_dimensions[1].height = 20
    ws_dash.row_dimensions[2].height = 28
    ws_dash.row_dimensions[3].height = 20
    ws_dash.row_dimensions[4].height = 10
    ws_dash.row_dimensions[5].height = 20
    ws_dash.row_dimensions[6].height = 24
    ws_dash.row_dimensions[7].height = 24
    ws_dash.row_dimensions[8].height = 20
    ws_dash.row_dimensions[9].height = 14
    ws_dash.row_dimensions[10].height = 26

    # Tarjetas KPI C-Level (Filas 5 a 8)
    # Card 1: B5:C8 -> Duración Esperada μ
    ws_dash.merge_cells("B5:C5")
    ws_dash["B5"] = "DURACIÓN ESPERADA (μ_total)" if lang == 'ES' else "EXPECTED DURATION (μ_total)"
    ws_dash["B5"].font = font_card_title
    ws_dash["B5"].fill = fill_card
    ws_dash["B5"].alignment = align_center

    ws_dash.merge_cells("B6:C7")
    ws_dash["B6"] = f"='{t_crit}'!C7"
    ws_dash["B6"].font = Font(name=FONT_NAME, size=20, bold=True, color=COLOR_BLUE_ACCENT)
    ws_dash["B6"].fill = fill_card
    ws_dash["B6"].alignment = align_center
    ws_dash["B6"].number_format = "0.0"

    ws_dash.merge_cells("B8:C8")
    ws_dash["B8"] = "días hábiles en camino crítico" if lang == 'ES' else "business days on critical path"
    ws_dash["B8"].font = font_card_sub
    ws_dash["B8"].fill = fill_card
    ws_dash["B8"].alignment = align_center

    # Card 2: D5:E8 -> Desviación Estándar σ
    ws_dash.merge_cells("D5:E5")
    ws_dash["D5"] = "DESVIACIÓN ESTÁNDAR (σ_total)" if lang == 'ES' else "STANDARD DEVIATION (σ_total)"
    ws_dash["D5"].font = font_card_title
    ws_dash["D5"].fill = fill_card
    ws_dash["D5"].alignment = align_center

    ws_dash.merge_cells("D6:E7")
    ws_dash["D6"] = f"='{t_crit}'!C9"
    ws_dash["D6"].font = Font(name=FONT_NAME, size=20, bold=True, color=COLOR_NAVY_DARK)
    ws_dash["D6"].fill = fill_card
    ws_dash["D6"].alignment = align_center
    ws_dash["D6"].number_format = "0.00"

    ws_dash.merge_cells("D8:E8")
    ws_dash["D8"] = "dispersión cuadrática CLT" if lang == 'ES' else "CLT quadratic dispersion"
    ws_dash["D8"].font = font_card_sub
    ws_dash["D8"].fill = fill_card
    ws_dash["D8"].alignment = align_center

    # Card 3: F5:G8 -> Intervalo de Certidumbre 95%
    ws_dash.merge_cells("F5:G5")
    ws_dash["F5"] = "INTERVALO DE CONFIANZA 95%" if lang == 'ES' else "95% CONFIDENCE INTERVAL"
    ws_dash["F5"].font = font_card_title
    ws_dash["F5"].fill = fill_card
    ws_dash["F5"].alignment = align_center

    ws_dash.merge_cells("F6:G7")
    ws_dash["F6"] = f'=ROUND(\'{t_crit}\'!C12, 1) & " a " & ROUND(\'{t_crit}\'!E12, 1)' if lang == 'ES' else f'=ROUND(\'{t_crit}\'!C12, 1) & " to " & ROUND(\'{t_crit}\'!E12, 1)'
    ws_dash["F6"].font = Font(name=FONT_NAME, size=13, bold=True, color=COLOR_NAVY_DARK)
    ws_dash["F6"].fill = fill_card
    ws_dash["F6"].alignment = align_center

    ws_dash.merge_cells("F8:G8")
    ws_dash["F8"] = "rango estadístico [μ ± 2σ]" if lang == 'ES' else "statistical range [μ ± 2σ]"
    ws_dash["F8"].font = font_card_sub
    ws_dash["F8"].fill = fill_card
    ws_dash["F8"].alignment = align_center

    # Card 4: H5:I8 -> Colchón al 90%
    ws_dash.merge_cells("H5:I5")
    ws_dash["H5"] = "COLCHÓN P90 CONSEJO" if lang == 'ES' else "BOARD P90 BUFFER REQUIRED"
    ws_dash["H5"].font = font_card_title
    ws_dash["H5"].fill = fill_green
    ws_dash["H5"].alignment = align_center

    ws_dash.merge_cells("H6:I7")
    ws_dash["H6"] = f"='{t_crit}'!D23"
    ws_dash["H6"].font = Font(name=FONT_NAME, size=20, bold=True, color=COLOR_GREEN_TEXT)
    ws_dash["H6"].fill = fill_green
    ws_dash["H6"].alignment = align_center
    ws_dash["H6"].number_format = "+0.0;-0.0;0.0"

    ws_dash.merge_cells("H8:I8")
    ws_dash["H8"] = "días extra para 90% certidumbre" if lang == 'ES' else "extra days for 90% certainty"
    ws_dash["H8"].font = font_card_sub
    ws_dash["H8"].fill = fill_green
    ws_dash["H8"].alignment = align_center

    # Aplicar fondos y bordes en las tarjetas
    for r in range(5, 9):
        for col_c in ["B", "C", "D", "E", "F", "G"]:
            ws_dash[f"{col_c}{r}"].fill = fill_card
            ws_dash[f"{col_c}{r}"].border = border_cell
        for col_c in ["H", "I"]:
            ws_dash[f"{col_c}{r}"].fill = fill_green
            ws_dash[f"{col_c}{r}"].border = border_cell

    # Calculadora Interactiva de Fecha Límite / Target Date (Filas 10 a 20)
    ws_dash["B10"] = "CALCULADORA INTERACTIVA DE FECHA COMPROMISO DIRECTIVA" if lang == 'ES' else "INTERACTIVE EXECUTIVE COMMITMENT DATE CALCULATOR"
    ws_dash["B10"].font = font_section

    calc_rows_es = [
        ("B11", "Fecha Límite / Días Objetivo Comprometidos (Td):", 188.0, "días hábiles", True, "0.0"),
        ("B12", "Duración Esperada del Modelo (μ_total):", f"='{t_crit}'!$C$7", "días hábiles", False, "0.0"),
        ("B13", "Desviación Estándar Combinada (σ_total):", f"='{t_crit}'!$C$9", "días hábiles", False, "0.00"),
        ("B14", "Z-Score Calculado [ (Td - μ) / σ ]:", f"=(C11 - C12) / C13", "valor estándar Z", False, "0.00"),
        ("B15", "Probabilidad de Éxito P(T ≤ Td):", f"=NORMDIST(C11, C12, C13, TRUE)", "% certidumbre", False, "0.0%"),
        ("B16", "Semáforo & Diagnóstico de Viabilidad:", f'=IF(C15>=0.85, "ALTA VIABILIDAD (>85%)", IF(C15>=0.6, "RIESGO MODERADO (60%-85%)", "INVIABLE SIN CONTINGENCIA (<60%)"))', "", False, "@"),
        ("B17", "Holgura (+) o Déficit (-) respecto a Media:", f"=C11 - C12", "días de margen", False, "+0.0;-0.0;0.0"),
        ("B18", "Días de Buffer Adicionales para 90% Certidumbre:", f"=MAX(0, '{t_crit}'!$C$23 - C11)", "días requeridos", False, "0.0"),
        ("B19", "Recomendación para el Comité de Dirección:", f'=IF(C15>=0.85, "Compromiso viable. Formalizar buffer en el contrato.", IF(C15>=0.6, "Incertidumbre elevada. Se requiere Fast-Tracking o colchón.", "VETO TÉCNICO. No firmar sin plan de Crashing o extensión formal."))', "", False, "@")
    ]
    calc_rows_en = [
        ("B11", "Committed Target Days / Deadline (Td):", 188.0, "business days", True, "0.0"),
        ("B12", "Model Expected Duration (μ_total):", f"='{t_crit}'!$C$7", "business days", False, "0.0"),
        ("B13", "Combined Standard Deviation (σ_total):", f"='{t_crit}'!$C$9", "business days", False, "0.00"),
        ("B14", "Calculated Z-Score [ (Td - μ) / σ ]:", f"=(C11 - C12) / C13", "standard Z value", False, "0.00"),
        ("B15", "Probability of Success P(T ≤ Td):", f"=NORMDIST(C11, C12, C13, TRUE)", "% certainty", False, "0.0%"),
        ("B16", "Feasibility Status & Traffic Light:", f'=IF(C15>=0.85, "HIGH FEASIBILITY (>85%)", IF(C15>=0.6, "MODERATE RISK (60%-85%)", "UNFEASIBLE W/O CONTINGENCY (<60%)"))', "", False, "@"),
        ("B17", "Slack (+) or Deficit (-) vs Expected Mean:", f"=C11 - C12", "margin days", False, "+0.0;-0.0;0.0"),
        ("B18", "Additional Buffer Days Needed for 90% Certainty:", f"=MAX(0, '{t_crit}'!$C$23 - C11)", "required days", False, "0.0"),
        ("B19", "Executive Governance Recommendation:", f'=IF(C15>=0.85, "Viable commitment. Formalize contingency buffer in contract.", IF(C15>=0.6, "High uncertainty. Fast-tracking or budget buffer required.", "TECHNICAL VETO. Do not sign without crashing or formal extension."))', "", False, "@")
    ]
    c_rows = calc_rows_es if lang == 'ES' else calc_rows_en
    input_ranges_dash = []

    for item in c_rows:
        cid, lbl, val, unit, is_inp, num_fmt = item
        rn = int(cid[1:])
        if rn == 16:
            ws_dash.row_dimensions[rn].height = 28
        elif rn == 18:
            ws_dash.row_dimensions[rn].height = 26
        elif rn == 19:
            ws_dash.row_dimensions[rn].height = 38
        else:
            ws_dash.row_dimensions[rn].height = 24

        ws_dash[cid] = lbl
        ws_dash[cid].font = font_body_bold
        ws_dash[cid].alignment = Alignment(horizontal="left", vertical="center")
        ws_dash[cid].border = border_cell

        ws_dash[f"C{rn}"] = val
        ws_dash[f"C{rn}"].font = font_input if is_inp else (Font(name=FONT_NAME, size=11, bold=True, color=COLOR_RED_TEXT if rn == 15 else COLOR_NAVY_DARK))
        ws_dash[f"C{rn}"].alignment = align_center
        ws_dash[f"C{rn}"].number_format = num_fmt
        ws_dash[f"C{rn}"].border = border_cell

        if is_inp:
            ws_dash[f"C{rn}"].fill = fill_input
            input_ranges_dash.append(f"C{rn}")
        elif rn == 15:
            ws_dash[f"C{rn}"].fill = fill_red
        elif rn == 16:
            ws_dash[f"C{rn}"].fill = fill_amber
            ws_dash[f"C{rn}"].font = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_AMBER_TEXT)
        else:
            ws_dash[f"C{rn}"].fill = fill_formula

        ws_dash[f"D{rn}"] = unit
        ws_dash[f"D{rn}"].font = Font(name=FONT_NAME, size=8.5, italic=True, color="64748B")
        ws_dash[f"D{rn}"].alignment = align_left
        ws_dash[f"D{rn}"].border = border_cell

    # Combinar celda de diagnóstico de viabilidad (C16:D16)
    ws_dash.merge_cells("C16:D16")
    ws_dash["C16"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_dash["D16"].fill = fill_amber
    ws_dash["D16"].border = border_cell

    # Combinar celda de recomendación para que quepa el texto (C19:E19)
    ws_dash.merge_cells("C19:E19")
    ws_dash["C19"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_dash["C19"].font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_NAVY_DARK)
    for col_m in ["C", "D", "E"]:
        ws_dash[f"{col_m}19"].border = border_cell
        ws_dash[f"{col_m}19"].fill = fill_formula

    # Tabla de Datos para la Curva de Densidad de Probabilidad (Gaussiana / Bell Curve)
    # Columna F: Días (x), Columna G: Densidad f(x), Columna H: Probabilidad Acumulada F(x)
    ws_dash["F10"] = "CURVA DE DENSIDAD DE PROBABILIDAD (GAUSSIANA CLT)" if lang == 'ES' else "PROBABILITY DENSITY FUNCTION (CLT GAUSSIAN)"
    ws_dash["F10"].font = font_section

    ws_dash["F11"] = "Días (x)" if lang == 'ES' else "Days (x)"
    ws_dash["G11"] = "Densidad f(x)" if lang == 'ES' else "Density f(x)"
    ws_dash["H11"] = "Probabilidad F(x)" if lang == 'ES' else "Cumulative F(x)"

    for col_g in ["F", "G", "H"]:
        ws_dash[f"{col_g}11"].font = font_tbl_header
        ws_dash[f"{col_g}11"].fill = fill_med
        ws_dash[f"{col_g}11"].alignment = align_center
        ws_dash[f"{col_g}11"].border = border_cell

    # Generar 25 puntos distribuidos en [μ - 3.5σ, μ + 3.5σ]
    x_str_list = []
    y_val_list = []
    mu_base = 192.5
    sigma_base = 10.14

    for idx_pt in range(25):
        row_curve = 12 + idx_pt
        ws_dash.row_dimensions[row_curve].height = 20

        # x = (μ - 3.5σ) + idx_pt * (7σ / 24)
        offset_k = -3.5 + idx_pt * (7.0 / 24.0)
        x_calc = round(mu_base + offset_k * sigma_base, 1)
        y_calc = (1.0 / (sigma_base * math.sqrt(2.0 * math.pi))) * math.exp(-0.5 * (offset_k ** 2))
        x_str_list.append(f"{x_calc:.1f}")
        y_val_list.append(round(y_calc, 5))

        ws_dash[f"F{row_curve}"] = f"=ROUND('{t_crit}'!$C$7 + ({offset_k:.4f}) * '{t_crit}'!$C$9, 1)"
        ws_dash[f"F{row_curve}"].font = font_body
        ws_dash[f"F{row_curve}"].alignment = align_center
        ws_dash[f"F{row_curve}"].number_format = "0.0"
        ws_dash[f"F{row_curve}"].fill = fill_formula
        ws_dash[f"F{row_curve}"].border = border_cell

        # f(x) = NORMDIST(x, μ, σ, FALSE)
        ws_dash[f"G{row_curve}"] = f"=NORMDIST(F{row_curve}, '{t_crit}'!$C$7, '{t_crit}'!$C$9, FALSE)"
        ws_dash[f"G{row_curve}"].font = font_body
        ws_dash[f"G{row_curve}"].alignment = align_center
        ws_dash[f"G{row_curve}"].number_format = "0.0000"
        ws_dash[f"G{row_curve}"].fill = fill_formula
        ws_dash[f"G{row_curve}"].border = border_cell

        # F(x) = NORMDIST(x, μ, σ, TRUE)
        ws_dash[f"H{row_curve}"] = f"=NORMDIST(F{row_curve}, '{t_crit}'!$C$7, '{t_crit}'!$C$9, TRUE)"
        ws_dash[f"H{row_curve}"].font = font_body
        ws_dash[f"H{row_curve}"].alignment = align_center
        ws_dash[f"H{row_curve}"].number_format = "0.0%"
        ws_dash[f"H{row_curve}"].fill = fill_formula
        ws_dash[f"H{row_curve}"].border = border_cell

    # Gráfico de Área / Densidad en Tab 1 (Columna J a Q, filas 10 a 30)
    chart = AreaChart()
    chart.title = "Distribución de Probabilidad del Cronograma (Curva de Campana)" if lang == 'ES' else "Schedule Probability Distribution (Bell Curve)"
    chart.style = 10
    chart.width = 20
    chart.height = 12

    data_ref = Reference(ws_dash, min_col=7, min_row=11, max_row=36)
    chart.add_data(data_ref, titles_from_data=True)

    # Configuración estricta de categorías (StrRef) y valores (NumRef) con caches
    cat_formula = f"'{ws_dash.title}'!$F$12:$F$36"
    sdata = StrData(ptCount=len(x_str_list), pt=[StrVal(idx=i, v=v) for i, v in enumerate(x_str_list)])
    sref = StrRef(f=cat_formula, strCache=sdata)

    val_formula = f"'{ws_dash.title}'!$G$12:$G$36"
    ndata = NumData(ptCount=len(y_val_list), pt=[NumVal(idx=i, v=str(v)) for i, v in enumerate(y_val_list)])
    nref = NumRef(f=val_formula, numCache=ndata)

    for s in chart.series:
        s.cat = AxDataSource(strRef=sref)
        s.val = NumDataSource(numRef=nref)
        s.graphicalProperties.solidFill = "2563EB"

    # Configuración explícita de ejes y marcas de graduación visibles
    chart.x_axis.title = "Duración del Proyecto (Días Hábiles)" if lang == 'ES' else "Project Duration (Business Days)"
    chart.x_axis.axPos = "b"
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.delete = False
    chart.x_axis.majorTickMark = "out"
    chart.x_axis.majorGridlines = ChartLines()
    chart.x_axis.tickLblSkip = 3

    chart.y_axis.title = "Densidad de Probabilidad f(x)" if lang == 'ES' else "Probability Density f(x)"
    chart.y_axis.axPos = "l"
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.delete = False
    chart.y_axis.majorTickMark = "out"
    chart.y_axis.majorGridlines = ChartLines()
    chart.y_axis.number_format = "0.000"

    chart.legend = None

    ws_dash.add_chart(chart, "J10")

    # Anchos de columna en Tab 1
    col_widths_dash = {
        "A": 4, "B": 52, "C": 26, "D": 22, "E": 4,
        "F": 16, "G": 18, "H": 20, "I": 4, "J": 16, "K": 16,
        "L": 16, "M": 16, "N": 16, "O": 16, "P": 16, "Q": 16
    }
    for col, width in col_widths_dash.items():
        ws_dash.column_dimensions[col].width = width

    # -------------------------------------------------------------------------
    # APLICACIÓN DE PROTECCIÓN DE HOJAS ECMA-376 (CONTRASEÑA "Datalaria2026")
    # -------------------------------------------------------------------------
    finalize_protection(ws_dash, input_ranges=[input_ranges_dash])
    finalize_protection(ws_wbs, input_ranges=[input_ranges_wbs])
    finalize_protection(ws_crit)
    finalize_protection(ws_buf, input_ranges=[input_ranges_buf])

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode = 'auto'

    return wb


def main():
    base_dir = "packages/Suite_03_Control_Operativo/03_Estimacion_PERT_3_Puntos"
    dir_es = os.path.join(base_dir, "[ES]_Estimacion_PERT_3_Puntos")
    dir_en = os.path.join(base_dir, "[EN]_3Point_PERT_Estimation")

    os.makedirs(dir_es, exist_ok=True)
    os.makedirs(dir_en, exist_ok=True)

    file_es = os.path.join(dir_es, "Estimacion_PERT_Estocastica_ES.xlsx")
    file_en = os.path.join(dir_en, "Stochastic_PERT_Estimation_EN.xlsx")

    print("[1/2] Generando modelo Excel en Español...")
    wb_es = create_pert_workbook(lang='ES')
    wb_es.save(file_es)
    print(f" -> Guardado exitosamente: {file_es}")

    print("[2/2] Generando modelo Excel en Inglés...")
    wb_en = create_pert_workbook(lang='EN')
    wb_en.save(file_en)
    print(f" -> Guardado exitosamente: {file_en}")

    print("Proceso completado con éxito.")


if __name__ == "__main__":
    main()
