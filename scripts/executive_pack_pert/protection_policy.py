#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
protection_policy.py
====================
Módulo de política de protección de celdas y seguridad OpenXML ECMA-376
para el Executive Decision Pack Estimación PERT de 3 Puntos Estocástica (Datalaria.com).

Reglas de protección:
- Celdas de entrada (usuario): locked=False, fondo blanco #FFFFFF.
- Celdas con fórmulas: locked=True, fondo #F1F5F9.
- Celdas decorativas/encabezados: locked=True, fondo corporativo preservado.
- Contraseña oficial: "Datalaria2026"
- ECMA-376: selectUnlockedCells=False, selectLockedCells=False
  (En XML, estos atributos son restricciones booleanas: False significa permitido).
"""

from openpyxl.styles import Protection, PatternFill
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.formula import ArrayFormula

PASSWORD_PROTECT = "Datalaria2026"
FILL_INPUT = PatternFill(fill_type="solid", fgColor="FFFFFF")
FILL_FORMULA = PatternFill(fill_type="solid", fgColor="F1F5F9")


def is_formula(cell):
    """Verifica si una celda contiene una fórmula de Excel."""
    v = cell.value
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def apply_sheet_protection(ws):
    """
    Aplica protección de hoja según el estándar ECMA-376 OpenXML.
    selectUnlockedCells = False (permite interactuar y editar celdas desbloqueadas con total normalidad)
    selectLockedCells = False (permite seleccionar celdas bloqueadas para lectura transparente)
    """
    ws.protection.set_password(PASSWORD_PROTECT)
    ws.protection.sheet = True
    ws.protection.objects = True
    ws.protection.scenarios = True
    ws.protection.selectUnlockedCells = False
    ws.protection.selectLockedCells = False
    ws.protection.insertRows = False
    ws.protection.deleteRows = False
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


def finalize_protection(ws, input_ranges=None, formula_ranges=None, decorative_cells=None):
    """
    Asegura que:
    1. Celdas de entrada especificadas tengan locked=False y fondo blanco.
    2. Celdas de fórmula tengan locked=True y fondo gris claro #F1F5F9.
    3. Celdas de cabecera/decorativas mantengan su estilo y locked=True.
    4. Se active la protección con contraseña oficial.
    """
    dec_set = set(decorative_cells) if decorative_cells else set()

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            coord = cell.coordinate
            if coord in dec_set:
                cell.protection = Protection(locked=True, hidden=False)
            elif is_formula(cell):
                cell.protection = Protection(locked=True, hidden=False)
                cell.fill = FILL_FORMULA
            elif input_ranges and any(coord in r for r in input_ranges):
                cell.protection = Protection(locked=False, hidden=False)
                cell.fill = FILL_INPUT

    apply_sheet_protection(ws)
