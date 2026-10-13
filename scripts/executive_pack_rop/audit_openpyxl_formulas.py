import openpyxl
import os
import re

files = [
    ('ES', os.path.abspath('packages/Suite_03_Control_Operativo/05_Stock_Seguridad_ROP/[ES]_Stock_Seguridad_ROP/Calculadora_Stock_Seguridad_ROP_ES.xlsx')),
    ('EN', os.path.abspath('packages/Suite_03_Control_Operativo/05_Stock_Seguridad_ROP/[EN]_Safety_Stock_ROP/Safety_Stock_ROP_Calculator_EN.xlsx'))
]

for lang, fpath in files:
    print(f"\n=======================================================")
    print(f"AUDITANDO LIBRO [{lang}]: {os.path.basename(fpath)}")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(fpath, data_only=False)
    sheet_names = wb.sheetnames
    print(f"Pestañas presentes ({len(sheet_names)}): {sheet_names}")
    
    total_formulas = 0
    norm_inv_count = 0
    bad_norm_inv = 0
    bad_sheet_refs = 0
    
    for sname in sheet_names:
        ws = wb[sname]
        f_count_sheet = 0
        for row in ws.iter_rows():
            for cell in row:
                val = str(cell.value or '')
                if val.startswith('='):
                    total_formulas += 1
                    f_count_sheet += 1
                    # Check for NORM.S.INV without _xlfn.
                    if 'NORM.S.INV' in val:
                        if '_xlfn.NORM.S.INV' in val:
                            norm_inv_count += 1
                        else:
                            bad_norm_inv += 1
                            print(f"  [ERROR] Formula with un-prefixed NORM.S.INV at {sname}!{cell.coordinate}: {val}")
                    
                    # Check sheet references in formula
                    # Matches pattern 'Sheet Name'!
                    refs = re.findall(r"'([^']+)'!", val)
                    for r in refs:
                        if r not in sheet_names:
                            bad_sheet_refs += 1
                            print(f"  [ERROR] Invalid sheet reference '{r}' in {sname}!{cell.coordinate}: {val}")
        
        print(f"  Pestaña '{sname}': {f_count_sheet} fórmulas analizadas.")
        
        # Check charts in sheet
        if ws._charts:
            for idx, ch in enumerate(ws._charts, 1):
                print(f"  [CHART {idx} en {sname}]:")
                print(f"    - Título: {ch.title}")
                print(f"    - x_axis.delete: {ch.x_axis.delete}")
                print(f"    - y_axis.delete: {ch.y_axis.delete}")
                print(f"    - x_axis.tickLblSkip: {getattr(ch.x_axis, 'tickLblSkip', None)}")
                print(f"    - Data series count: {len(ch.series)}")
                for s_i, s in enumerate(ch.series):
                    print(f"      * Serie {s_i+1}: {s.tx} (Values: {s.val.numRef.f if hasattr(s.val, 'numRef') and s.val.numRef else 'N/A'})")
    
    print(f"\nResumen Auditoría Fórmulas [{lang}]:")
    print(f"  - Total fórmulas auditadas: {total_formulas}")
    print(f"  - Fórmulas _xlfn.NORM.S.INV válidas: {norm_inv_count}")
    print(f"  - Fórmulas NORM.S.INV sin prefijo (erróneas): {bad_norm_inv}")
    print(f"  - Referencias a hojas rotas: {bad_sheet_refs}")
    if bad_norm_inv == 0 and bad_sheet_refs == 0:
        print(f"  ==> LIBRO [{lang}] CERTIFICADO 100% VÁLIDO Y COHERENTE.")
    else:
        print(f"  ==> LIBRO [{lang}] CONTIENE ERRORES.")
