#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_excel.py
=================
Genera los modelos analíticos en Excel (.xlsx) de grado Consejo de Administración para el Executive Decision Pack:
1. packages/Suite_03_Control_Operativo/04_Cuadro_EVM_Valor_Ganado/[ES]_Cuadro_EVM_Valor_Ganado/Cuadro_Mando_EVM_Datalaria_ES.xlsx
2. packages/Suite_03_Control_Operativo/04_Cuadro_EVM_Valor_Ganado/[EN]_Earned_Value_Management_EVM/EVM_Dashboard_Datalaria_EN.xlsx

Estructura de 4 Pestañas Funcionales:
- Pestaña 1: "Dashboard Ejecutivo EVM" / "Executive EVM Dashboard"
- Pestaña 2: "Control WBS & Entregables" / "WBS & Deliverables Tracking"
- Pestaña 3: "Serie Temporal & Proyecciones" / "Time Series & EAC Forecasting"
- Pestaña 4: "Plan de Acción & Recuperación" / "Action & Recovery Plan"

Estándar de Seguridad y Compatibilidad OpenXML ECMA-376:
- Celdas de entrada (usuario): locked=False, fondo blanco #FFFFFF.
- Celdas de fórmulas y títulos: locked=True, fondo #F1F5F9 o corporativo.
- Contraseña oficial interna: "Datalaria2026"
- ECMA-376: selectUnlockedCells = False, selectLockedCells = False.
- Curva S nativa generada mediante openpyxl LineChart con trazado multi-serie y ejes completamente configurados.
- Alturas de fila y anchos de columna generosos para garantizar legibilidad ejecutiva sin textos cortados ni solapes.
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrRef, StrData, StrVal

# Importar política de protección ECMA-376
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from protection_policy import finalize_protection

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# --- PALETA CORPORATIVA DATALARIA ---
COLOR_NAVY_DARK = "0F172A"       # Slate 900
COLOR_NAVY_MED = "1E293B"        # Slate 800
COLOR_NAVY_LIGHT = "334155"      # Slate 700
COLOR_BLUE_ACCENT = "2563EB"     # Blue 600
COLOR_BLUE_LIGHT = "EFF6FF"      # Blue 50
COLOR_CYAN_ACCENT = "0284C7"     # Sky 600
COLOR_PURPLE_ACCENT = "7C3AED"   # Purple 600 (EAC Forecast)

COLOR_GREEN_BG = "D1FAE5"        # Emerald 100
COLOR_GREEN_TEXT = "065F46"      # Emerald 800
COLOR_GREEN_BORDER = "10B981"    # Emerald 500

COLOR_AMBER_BG = "FEF3C7"        # Amber 100
COLOR_AMBER_TEXT = "92400E"      # Amber 800
COLOR_AMBER_BORDER = "F59E0B"    # Amber 500

COLOR_RED_BG = "FEE2E2"          # Red 100
COLOR_RED_TEXT = "991B1B"        # Red 800
COLOR_RED_BORDER = "DC2626"      # Red 600

COLOR_BG_CARD = "F8FAFC"         # Slate 50
COLOR_BG_INPUT = "FFFFFF"        # Blanco puro para inputs
COLOR_BG_FORMULA = "F1F5F9"      # Slate 100 para fórmulas
COLOR_BORDER_LIGHT = "CBD5E1"    # Slate 300
COLOR_BORDER_DARK = "64748B"     # Slate 500

FONT_NAME = "Segoe UI"

# Estilos reutilizables
font_title = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
font_subtitle = Font(name=FONT_NAME, size=9.5, italic=True, color="94A3B8")
font_section = Font(name=FONT_NAME, size=10.5, bold=True, color=COLOR_NAVY_DARK)
font_tbl_header = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
font_tbl_subheader = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
font_card_title = Font(name=FONT_NAME, size=8.5, bold=True, color="64748B")
font_card_val = Font(name=FONT_NAME, size=15, bold=True, color=COLOR_NAVY_DARK)
font_card_sub = Font(name=FONT_NAME, size=8, italic=True, color="64748B")
font_body = Font(name=FONT_NAME, size=9.5, color="1E293B")
font_body_bold = Font(name=FONT_NAME, size=9.5, bold=True, color="1E293B")
font_input = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_BLUE_ACCENT)

fill_dark = PatternFill(fill_type="solid", fgColor=COLOR_NAVY_DARK)
fill_med = PatternFill(fill_type="solid", fgColor=COLOR_NAVY_MED)
fill_blue_accent = PatternFill(fill_type="solid", fgColor=COLOR_BLUE_ACCENT)
fill_blue_card = PatternFill(fill_type="solid", fgColor=COLOR_BLUE_LIGHT)
fill_card = PatternFill(fill_type="solid", fgColor=COLOR_BG_CARD)
fill_formula = PatternFill(fill_type="solid", fgColor=COLOR_BG_FORMULA)
fill_input = PatternFill(fill_type="solid", fgColor=COLOR_BG_INPUT)

fill_green = PatternFill(fill_type="solid", fgColor=COLOR_GREEN_BG)
fill_amber = PatternFill(fill_type="solid", fgColor=COLOR_AMBER_BG)
fill_red = PatternFill(fill_type="solid", fgColor=COLOR_RED_BG)

thin_border_side = Side(style="thin", color=COLOR_BORDER_LIGHT)
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_thick_side = Side(style="medium", color=COLOR_BLUE_ACCENT)
border_card = Border(left=border_thick_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_top_total = Border(top=Side(style="thin", color=COLOR_NAVY_DARK), bottom=Side(style="double", color=COLOR_NAVY_DARK))

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)


def get_wbs_data(lang='ES'):
    """
    Datos estructurados de 25 entregables en 5 fases calibrados con exactitud quirúrgica:
    - Presupuesto Total: BAC = 1.200.000 € / $1,200,000
    - A corte Mes 6:
      * PV = 600.000 €
      * EV = 516.000 €
      * AC = 600.000 €
      * CPI = 0.86, SPI = 0.86, CV = -84.000 €, SV = -84.000 €
    """
    if lang == 'ES':
        phases = [
            ("Fase 1: Arquitectura, Gobierno & Ciberseguridad", [
                ("1.1", "Definición de Arquitectura Cloud & Microservicios", 40000, 1.00, 40000, 41000),
                ("1.2", "Modelado Conceptual de Datos & Seguridad IAM", 30000, 1.00, 30000, 30500),
                ("1.3", "Especificación de APIs OpenAPI & API Gateway", 35000, 1.00, 35000, 36000),
                ("1.4", "Marco de Cumplimiento Regulatorio & GDPR / DORA", 30000, 1.00, 30000, 30000),
                ("1.5", "Aprovisionamiento de Infraestructura IaC Terraform", 45000, 1.00, 45000, 46000),
            ]),
            ("Fase 2: Desarrollo Core & Motor Transaccional", [
                ("2.1", "Desarrollo del Motor Transaccional de Conciliación", 85000, 0.85, 85000, 98000),
                ("2.2", "Implementación de Autenticación OAuth2 / OIDC", 50000, 0.90, 50000, 52000),
                ("2.3", "Módulo de Procesamiento de Pagos & Pasarelas", 75000, 0.70, 75000, 82000),
                ("2.4", "Motor de Reglas de Negocio & Enrutamiento", 65000, 0.65, 65000, 68000),
                ("2.5", "Sistema de Auditoría & Trazabilidad de Logs", 35000, 0.80, 35000, 36500),
            ]),
            ("Fase 3: Integraciones, APIs & Conectores", [
                ("3.1", "Conector Bidireccional con ERP Core (SAP / Oracle)", 80000, 0.40, 35000, 28000),
                ("3.2", "Integración con Pasarelas Bancarias & Red SWIFT", 60000, 0.40, 25000, 20000),
                ("3.3", "Conector con CRM Salesforce & Servicio Clientes", 40000, 0.25, 15000, 10000),
                ("3.4", "Bus de Eventos Kafka & Sincronización Real-Time", 50000, 0.36, 20000, 12000),
                ("3.5", "Pipeline ETL de Reportería & Business Intelligence", 50000, 0.24, 15000, 10000),
            ]),
            ("Fase 4: Migración de Datos, QA & Certificación", [
                ("4.1", "Estrategia & Scripts de Migración de Datos Históricos", 60000, 0.00, 0, 0),
                ("4.2", "Banco de Pruebas Automatizadas Unitarias & E2E", 45000, 0.00, 0, 0),
                ("4.3", "Pruebas de Estrés, Carga & Rendimiento (JMeter)", 35000, 0.00, 0, 0),
                ("4.4", "Auditoría de Ciberseguridad & Pentest Certificado", 40000, 0.00, 0, 0),
                ("4.5", "Certificación UAT Formal con Usuarios de Negocio", 50000, 0.00, 0, 0),
            ]),
            ("Fase 5: Infraestructura Cloud, Despliegue & Cutover", [
                ("5.1", "Configuración de Clusters Kubernetes & Malla Istio", 45000, 0.00, 0, 0),
                ("5.2", "Pipeline CI/CD con Despliegue Blue-Green", 35000, 0.00, 0, 0),
                ("5.3", "Plan de Recuperación de Desastres (DRP) Multi-Región", 40000, 0.00, 0, 0),
                ("5.4", "Ensayo General de Migración (Dry-Run Cutover)", 40000, 0.00, 0, 0),
                ("5.5", "Pase a Producción Final & Hipercare Post Go-Live", 40000, 0.00, 0, 0),
            ]),
        ]
    else:
        phases = [
            ("Phase 1: Architecture, Governance & Cybersecurity", [
                ("1.1", "Cloud Architecture & Microservices Definition", 40000, 1.00, 40000, 41000),
                ("1.2", "Data Modeling & IAM Security Framework", 30000, 1.00, 30000, 30500),
                ("1.3", "OpenAPI Specification & Enterprise API Gateway", 35000, 1.00, 35000, 36000),
                ("1.4", "Regulatory Compliance Framework & GDPR / DORA", 30000, 1.00, 30000, 30000),
                ("1.5", "IaC Terraform Infrastructure Provisioning", 45000, 1.00, 45000, 46000),
            ]),
            ("Phase 2: Core Platform & Transaction Engine", [
                ("2.1", "Core Reconciliation & Transaction Engine Dev", 85000, 0.85, 85000, 98000),
                ("2.2", "OAuth2 / OIDC Enterprise Authentication Flow", 50000, 0.90, 50000, 52000),
                ("2.3", "Payment Processing Module & Gateway Routing", 75000, 0.70, 75000, 82000),
                ("2.4", "Business Rules Engine & Automated Routing", 65000, 0.65, 65000, 68000),
                ("2.5", "Audit Trail & Immutable Log Tracing System", 35000, 0.80, 35000, 36500),
            ]),
            ("Phase 3: Integrations, APIs & Connectors", [
                ("3.1", "Bidirectional Core ERP Connector (SAP / Oracle)", 80000, 0.40, 35000, 28000),
                ("3.2", "Banking Gateway & SWIFT Network Integration", 60000, 0.40, 25000, 20000),
                ("3.3", "CRM Salesforce & Customer Care Integration", 40000, 0.25, 15000, 10000),
                ("3.4", "Kafka Event Bus & Real-Time Sync Streams", 50000, 0.36, 20000, 12000),
                ("3.5", "ETL Data Pipeline & BI Executive Reporting", 50000, 0.24, 15000, 10000),
            ]),
            ("Phase 4: Data Migration, QA & Certification", [
                ("4.1", "Legacy Data Cleansing & Migration Scripts", 60000, 0.00, 0, 0),
                ("4.2", "Automated Unit & End-to-End Test Suite", 45000, 0.00, 0, 0),
                ("4.3", "Stress, Load & Performance Benchmarking (JMeter)", 35000, 0.00, 0, 0),
                ("4.4", "Cybersecurity Audit & Penetration Testing", 40000, 0.00, 0, 0),
                ("4.5", "Formal UAT Certification with Business Leads", 50000, 0.00, 0, 0),
            ]),
            ("Phase 5: Cloud Infra, Deployment & Cutover", [
                ("5.1", "Kubernetes Production Clusters & Istio Mesh", 45000, 0.00, 0, 0),
                ("5.2", "Zero-Downtime Blue-Green CI/CD Pipelines", 35000, 0.00, 0, 0),
                ("5.3", "Multi-Region Disaster Recovery Plan (DRP)", 40000, 0.00, 0, 0),
                ("5.4", "Dress Rehearsal & Cutover Simulation (Dry-Run)", 40000, 0.00, 0, 0),
                ("5.5", "Production Go-Live & Post-Launch Hypercare", 40000, 0.00, 0, 0),
            ]),
        ]
    return phases


def get_monthly_data(lang='ES'):
    """
    Serie temporal de 12 meses alineada exactamente con BAC = 1.200.000 €:
    - M01 a M06 suman exactamente PV = 600.000 €, EV = 516.000 €, AC = 600.000 €.
    - M07 a M12 completan el PV restante hasta 1.200.000 €.
    """
    return [
        ("M01", 60000, 58000, 62000),
        ("M02", 80000, 75000, 84000),
        ("M03", 100000, 92000, 108000),
        ("M04", 110000, 95000, 114000),
        ("M05", 120000, 98000, 116000),
        ("M06", 130000, 98000, 116000),  # Cutoff en M06: Cum PV=600k, Cum EV=516k, Cum AC=600k
        ("M07", 130000, "", ""),
        ("M08", 120000, "", ""),
        ("M09", 110000, "", ""),
        ("M10", 100000, "", ""),
        ("M11", 80000, "", ""),
        ("M12", 60000, "", ""),          # Cum PV total = 1.200.000 €
    ]


def build_tab1_dashboard(ws, lang='ES'):
    """Construye la Pestaña 1: Dashboard Ejecutivo EVM con alturas y anchos calibrados."""
    fmt_curr = '#,##0 "€"' if lang == 'ES' else '"$"#,##0'
    fmt_idx = '0.00"x"'
    fmt_pct = '0.0%'

    # Anchos de columna uniformes y generosos para evitar textos cortados
    col_widths = {
        'A': 3,
        'B': 16, 'C': 16, 'D': 16,
        'E': 16, 'F': 16, 'G': 16,
        'H': 16, 'I': 16, 'J': 16,
        'K': 16, 'L': 16, 'M': 16,
        'N': 16, 'O': 16, 'P': 16,
        'Q': 3
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Alturas de fila explícitas para evitar solapes verticales en cuadrantes, tarjetas y gráfico
    row_heights = {
        1: 12,
        2: 24, 3: 24,   # Banner principal
        4: 22,          # Metadatos
        5: 10,          # Espacio
        6: 20, 7: 32, 8: 18,   # KPI Row 1
        9: 10,                 # Espacio
        10: 20, 11: 32, 12: 18,# KPI Row 2
        13: 10,                # Espacio
        14: 20, 15: 32, 16: 18,# KPI Row 3
        17: 12,                # Espacio
        18: 26,                # Títulos de sección
        19: 25, 20: 25,        # Cuadrantes 1 y 2 (Total 50 pt)
        21: 25, 22: 25,        # Cuadrantes 3 y 4 (Total 50 pt)
        23: 22, 24: 22, 25: 22, 26: 22, # Diagnóstico ejecutivo (Total 88 pt)
        27: 22, 28: 22, 29: 22, 30: 22, 31: 22, 32: 22, 33: 22, 34: 22, 35: 22, 36: 22
    }
    for r, h in row_heights.items():
        ws.row_dimensions[r].height = h

    ws.views.sheetView[0].showGridLines = True

    # 1. Banner Principal
    ws.merge_cells('B2:P3')
    top_cell = ws['B2']
    title_text = "DATALARIA | CUADRO DE MANDO EJECUTIVO EVM (EARNED VALUE MANAGEMENT)" if lang == 'ES' else "DATALARIA | EXECUTIVE EVM DASHBOARD (EARNED VALUE MANAGEMENT)"
    sub_text = "Estándar ANSI/EIA-748 & PMBOK · Diagnóstico de Curva S, Varianzas & Proyecciones EAC C-Level" if lang == 'ES' else "ANSI/EIA-748 & PMBOK Standards · S-Curve Diagnostics, Cost/Schedule Variances & C-Level EAC"
    top_cell.value = f"{title_text}\n{sub_text}"
    top_cell.font = font_title
    top_cell.fill = fill_dark
    top_cell.alignment = align_center

    # 2. Metadatos de Corte
    ws['B4'] = "PROYECTO: Transformación Core Transaccional" if lang == 'ES' else "PROJECT: Core Transactional Modernization"
    ws['B4'].font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_NAVY_DARK)
    ws['K4'] = "FECHA DE CORTE: Mes 6 (50% Cronograma Transcurrido)" if lang == 'ES' else "CUTOFF DATE: Month 6 (50% Elapsed Schedule)"
    ws['K4'].font = Font(name=FONT_NAME, size=9.5, bold=True, color=COLOR_BLUE_ACCENT)

    # 3. FILA 1 DE KPI CARDS: Cifras Financieras Clave (Filas 6-8)
    ref_wbs = "'Control WBS & Entregables'" if lang == 'ES' else "'WBS & Deliverables Tracking'"
    kpi_blocks_1 = [
        ('B', 'D', "PRESUPUESTO TOTAL (BAC)" if lang == 'ES' else "BUDGET AT COMPLETION (BAC)",
         f"={ref_wbs}!E32",
         fmt_curr, "Línea base contractual aprobada" if lang == 'ES' else "Approved contractual baseline"),

        ('E', 'G', "VALOR PLANIFICADO (PV)" if lang == 'ES' else "PLANNED VALUE (PV)",
         f"={ref_wbs}!H32",
         fmt_curr, "Trabajo programado a corte" if lang == 'ES' else "Scheduled work at cutoff"),

        ('H', 'J', "VALOR GANADO FÍSICO (EV)" if lang == 'ES' else "EARNED VALUE (EV)",
         f"={ref_wbs}!J32",
         fmt_curr, "Valor del trabajo entregado" if lang == 'ES' else "Value of delivered scope"),

        ('K', 'M', "COSTE REAL INCURRIDO (AC)" if lang == 'ES' else "ACTUAL COST (AC)",
         f"={ref_wbs}!I32",
         fmt_curr, "Facturación & recursos gastados" if lang == 'ES' else "Total money spent to date"),

        ('N', 'P', "% AVANCE REAL" if lang == 'ES' else "% ACTUAL PROGRESS",
         "=H7/B7", fmt_pct, "EV / BAC acumulado" if lang == 'ES' else "Cumulative EV / BAC"),
    ]

    for c_start, c_end, title, formula, num_fmt, note in kpi_blocks_1:
        ws.merge_cells(f'{c_start}6:{c_end}6')
        ws.merge_cells(f'{c_start}7:{c_end}7')
        ws.merge_cells(f'{c_start}8:{c_end}8')

        c_t = ws[f'{c_start}6']
        c_v = ws[f'{c_start}7']
        c_n = ws[f'{c_start}8']

        c_t.value = title
        c_t.font = font_card_title
        c_t.fill = fill_card
        c_t.alignment = align_center

        c_v.value = formula
        c_v.font = font_card_val
        c_v.number_format = num_fmt
        c_v.fill = fill_card
        c_v.alignment = align_center

        c_n.value = note
        c_n.font = font_card_sub
        c_n.fill = fill_card
        c_n.alignment = align_center

        for row in range(6, 9):
            for col_idx in range(openpyxl.utils.column_index_from_string(c_start), openpyxl.utils.column_index_from_string(c_end) + 1):
                col_letter = get_column_letter(col_idx)
                cell = ws[f'{col_letter}{row}']
                cell.border = border_cell

    # 4. FILA 2 DE KPI CARDS: Índices de Desempeño y Varianzas (Filas 10-12)
    kpi_blocks_2 = [
        ('B', 'D', "ÍNDICE EFICIENCIA COSTE (CPI)" if lang == 'ES' else "COST PERFORMANCE INDEX (CPI)",
         f"={ref_wbs}!M32",
         fmt_idx, "EV / AC (<1.0 indica sobrecoste)" if lang == 'ES' else "EV / AC (<1.0 means over budget)"),

        ('E', 'G', "ÍNDICE EFICIENCIA PLAZO (SPI)" if lang == 'ES' else "SCHEDULE PERF. INDEX (SPI)",
         f"={ref_wbs}!N32",
         fmt_idx, "EV / PV (<1.0 indica retraso)" if lang == 'ES' else "EV / PV (<1.0 means delayed)"),

        ('H', 'J', "VARIANZA DE COSTE (CV)" if lang == 'ES' else "COST VARIANCE (CV)",
         f"={ref_wbs}!K32",
         fmt_curr, "EV - AC (Negativo = Sobrecoste)" if lang == 'ES' else "EV - AC (Negative = Overrun)"),

        ('K', 'M', "VARIANZA DE CRONOGRAMA (SV)" if lang == 'ES' else "SCHEDULE VARIANCE (SV)",
         f"={ref_wbs}!L32",
         fmt_curr, "EV - PV (Negativo = Retraso)" if lang == 'ES' else "EV - PV (Negative = Late)"),

        ('N', 'P', "DESVÍO EFECTIVO" if lang == 'ES' else "EFFICIENCY GAP",
         "=(B11-1)", fmt_pct, "Pérdida de poder adquisitivo" if lang == 'ES' else "Purchasing power loss"),
    ]

    for c_start, c_end, title, formula, num_fmt, note in kpi_blocks_2:
        ws.merge_cells(f'{c_start}10:{c_end}10')
        ws.merge_cells(f'{c_start}11:{c_end}11')
        ws.merge_cells(f'{c_start}12:{c_end}12')

        c_t = ws[f'{c_start}10']
        c_v = ws[f'{c_start}11']
        c_n = ws[f'{c_start}12']

        c_t.value = title
        c_t.font = font_card_title
        c_t.fill = fill_blue_card
        c_t.alignment = align_center

        c_v.value = formula
        c_v.font = Font(name=FONT_NAME, size=15, bold=True, color=COLOR_RED_TEXT if 'CPI' in title or 'SPI' in title or 'VARIANZA' in title else COLOR_NAVY_DARK)
        c_v.number_format = num_fmt
        c_v.fill = fill_blue_card
        c_v.alignment = align_center

        c_n.value = note
        c_n.font = font_card_sub
        c_n.fill = fill_blue_card
        c_n.alignment = align_center

        for row in range(10, 13):
            for col_idx in range(openpyxl.utils.column_index_from_string(c_start), openpyxl.utils.column_index_from_string(c_end) + 1):
                col_letter = get_column_letter(col_idx)
                cell = ws[f'{col_letter}{row}']
                cell.border = border_cell

    # 5. FILA 3 DE KPI CARDS: Proyecciones Estadísticas de Cierre (Filas 14-16)
    ref_ts = "'Serie Temporal & Proyecciones'" if lang == 'ES' else "'Time Series & EAC Forecasting'"
    kpi_blocks_3 = [
        ('B', 'D', "PREVISIÓN AL CIERRE (EAC 1 TÍPICO)" if lang == 'ES' else "ESTIMATE AT COMPLETION (EAC 1)",
         f"={ref_ts}!D23",
         fmt_curr, "BAC / CPI (Tendencia actual)" if lang == 'ES' else "BAC / CPI (Current trend)"),

        ('E', 'G', "DESVIACIÓN FINAL PREVISTA (VAC)" if lang == 'ES' else "VARIANCE AT COMPLETION (VAC)",
         f"={ref_ts}!D26",
         fmt_curr, "BAC - EAC (Déficit presupuestario)" if lang == 'ES' else "BAC - EAC (Budget deficit)"),

        ('H', 'J', "ÍNDICE REQUERIDO TCPI (BAC)" if lang == 'ES' else "TO-COMPLETE INDEX TCPI (BAC)",
         f"={ref_ts}!D28",
         fmt_idx, "(BAC-EV)/(BAC-AC) [>1.10 Inviable]" if lang == 'ES' else "(BAC-EV)/(BAC-AC) [>1.10 Unviable]"),

        ('K', 'M', "ÍNDICE REQUERIDO TCPI (EAC)" if lang == 'ES' else "TO-COMPLETE INDEX TCPI (EAC)",
         f"={ref_ts}!D29",
         fmt_idx, "(BAC-EV)/(EAC-AC) [Meta revisada]" if lang == 'ES' else "(BAC-EV)/(EAC-AC) [Revised target]"),

        ('N', 'P', "SEVERIDAD ALERTA" if lang == 'ES' else "ALERT STATUS",
         '=IF(H15>1.10, "CRÍTICA", "CONTROLADO")' if lang == 'ES' else '=IF(H15>1.10, "CRITICAL", "CONTROLLED")',
         '@', "Umbral de riesgo PMBOK" if lang == 'ES' else "PMBOK Risk Threshold"),
    ]

    for c_start, c_end, title, formula, num_fmt, note in kpi_blocks_3:
        ws.merge_cells(f'{c_start}14:{c_end}14')
        ws.merge_cells(f'{c_start}15:{c_end}15')
        ws.merge_cells(f'{c_start}16:{c_end}16')

        c_t = ws[f'{c_start}14']
        c_v = ws[f'{c_start}15']
        c_n = ws[f'{c_start}16']

        c_t.value = title
        c_t.font = font_card_title
        c_t.fill = fill_amber if 'VAC' in title or 'TCPI' in title else fill_card
        c_t.alignment = align_center

        c_v.value = formula
        c_v.font = Font(name=FONT_NAME, size=15, bold=True, color=COLOR_RED_TEXT if 'VAC' in title else COLOR_NAVY_DARK)
        c_v.number_format = num_fmt
        c_v.fill = fill_amber if 'VAC' in title or 'TCPI' in title else fill_card
        c_v.alignment = align_center

        c_n.value = note
        c_n.font = font_card_sub
        c_n.fill = fill_amber if 'VAC' in title or 'TCPI' in title else fill_card
        c_n.alignment = align_center

        for row in range(14, 17):
            for col_idx in range(openpyxl.utils.column_index_from_string(c_start), openpyxl.utils.column_index_from_string(c_end) + 1):
                col_letter = get_column_letter(col_idx)
                cell = ws[f'{col_letter}{row}']
                cell.border = border_cell

    # 6. SECCIÓN INFERIOR: SEMÁFORO 2x2 & DIAGNÓSTICO (Cols B a G) Y CURVA S (Cols I a O)
    ws.merge_cells('B18:G18')
    ws['B18'] = "MATRIZ DE SALUD OPERATIVA (CPI vs SPI) & GOBERNANZA" if lang == 'ES' else "OPERATING HEALTH MATRIX (CPI vs SPI) & GOVERNANCE"
    ws['B18'].font = font_section
    ws['B18'].fill = fill_formula
    ws['B18'].alignment = align_left

    # Cuadrantes 2x2 con filas separadas y altura suficiente (Rows 19-20 y 21-22)
    # Cuadrante 1: CPI >= 1 & SPI >= 1
    ws.merge_cells('B19:D20')
    ws['B19'] = ("CUADRANTE 1: SALUDABLE\nCPI ≥ 1.00  |  SPI ≥ 1.00\nEn plazo y dentro de presupuesto" if lang == 'ES'
                 else "QUADRANT 1: HEALTHY\nCPI ≥ 1.00  |  SPI ≥ 1.00\nOn time and within budget")
    ws['B19'].font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_GREEN_TEXT)
    ws['B19'].fill = fill_green
    ws['B19'].alignment = align_center

    # Cuadrante 2: CPI >= 1 & SPI < 1
    ws.merge_cells('E19:G20')
    ws['E19'] = ("CUADRANTE 2: ALERTA PLAZO\nCPI ≥ 1.00  |  SPI < 1.00\nBajo presupuesto pero con retraso" if lang == 'ES'
                 else "QUADRANT 2: SCHEDULE ALERT\nCPI ≥ 1.00  |  SPI < 1.00\nUnder budget but behind schedule")
    ws['E19'].font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_AMBER_TEXT)
    ws['E19'].fill = fill_amber
    ws['E19'].alignment = align_center

    # Cuadrante 3: CPI < 1 & SPI >= 1
    ws.merge_cells('B21:D22')
    ws['B21'] = ("CUADRANTE 3: ALERTA COSTE\nCPI < 1.00  |  SPI ≥ 1.00\nEn plazo pero con sobrecoste" if lang == 'ES'
                 else "QUADRANT 3: COST ALERT\nCPI < 1.00  |  SPI ≥ 1.00\nOn schedule but over budget")
    ws['B21'].font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_AMBER_TEXT)
    ws['B21'].fill = fill_amber
    ws['B21'].alignment = align_center

    # Cuadrante 4: CPI < 1 & SPI < 1 (CRISIS)
    ws.merge_cells('E21:G22')
    ws['E21'] = ("CUADRANTE 4: ZONA DE CRISIS [CRÍTICA]\nCPI < 1.00  |  SPI < 1.00\nSobrecoste activo y retraso severo" if lang == 'ES'
                 else "QUADRANT 4: CRISIS ZONE [CRITICAL]\nCPI < 1.00  |  SPI < 1.00\nOver budget and behind schedule")
    ws['E21'].font = Font(name=FONT_NAME, size=9, bold=True, color=COLOR_RED_TEXT)
    ws['E21'].fill = fill_red
    ws['E21'].alignment = align_center

    # Bordes de cuadrantes
    for r in range(19, 23):
        for c in range(2, 8):
            ws.cell(row=r, column=c).border = border_cell

    # Diagnóstico dinámico en texto (Filas 23 a 26 con 88 pt de altura total)
    ws.merge_cells('B23:G26')
    diag_cell = ws['B23']
    diag_formula = (
        '="DIAGNÓSTICO EJECUTIVO: El proyecto se encuentra actualmente en el " & '
        'IF(AND(B11>=1,E11>=1),"CUADRANTE 1 (SALUDABLE)",'
        'IF(AND(B11<1,E11>=1),"CUADRANTE 3 (ALERTA COSTE)",'
        'IF(AND(B11>=1,E11<1),"CUADRANTE 2 (ALERTA PLAZO)",'
        '"CUADRANTE 4 (CRISIS CRÍTICA)"))) & ". " & CHAR(10) & '
        '"• Eficiencia de Coste: Por cada euro gastado solo se generan " & TEXT(B11,"0.00") & " € de valor físico." & CHAR(10) & '
        '"• Retraso acumulado: Déficit de cronograma valorado en " & TEXT(K11,"#,##0 €") & ". " & CHAR(10) & '
        '"• Conclusión Board: Se requiere activación inmediata de Plan de Recuperación y autorización de techo EAC."'
        if lang == 'ES' else
        '="EXECUTIVE DIAGNOSIS: The project is currently positioned in " & '
        'IF(AND(B11>=1,E11>=1),"QUADRANT 1 (HEALTHY)",'
        'IF(AND(B11<1,E11>=1),"QUADRANT 3 (COST ALERT)",'
        'IF(AND(B11>=1,E11<1),"QUADRANT 2 (SCHEDULE ALERT)",'
        '"QUADRANT 4 (CRISIS ZONE)"))) & ". " & CHAR(10) & '
        '"• Cost Efficiency: For every dollar spent, only $" & TEXT(B11,"0.00") & " of earned value is generated." & CHAR(10) & '
        '"• Schedule Deficit: Accumulated delay valued at " & TEXT(K11,"$#,##0") & ". " & CHAR(10) & '
        '"• Board Takeaway: Immediate Recovery Plan activation and EAC revision approval required."'
    )
    diag_cell.value = diag_formula
    diag_cell.font = Font(name=FONT_NAME, size=9, italic=False, color=COLOR_NAVY_DARK)
    diag_cell.fill = fill_card
    diag_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for r in range(23, 27):
        for c in range(2, 8):
            ws.cell(row=r, column=c).border = border_cell

    # Título para el Gráfico de la Curva S (Cols H a P)
    ws.merge_cells('H18:P18')
    ws['H18'] = "CURVA S EJECUTIVA: PV vs EV vs AC & PROYECCIÓN EAC" if lang == 'ES' else "EXECUTIVE S-CURVE: PV vs EV vs AC & EAC FORECAST"
    ws['H18'].font = font_section
    ws['H18'].fill = fill_formula
    ws['H18'].alignment = align_left


def build_tab2_wbs(ws, lang='ES'):
    """Construye la Pestaña 2: Control WBS & Entregables calibrada a BAC = 1.200.000 €."""
    fmt_curr = '#,##0 "€"' if lang == 'ES' else '"$"#,##0'
    fmt_idx = '0.00"x"'
    fmt_pct = '0.0%'

    col_widths = {
        'A': 3,
        'B': 8,   # WBS
        'C': 34,  # Fase
        'D': 44,  # Entregable
        'E': 17,  # BAC
        'F': 12,  # Peso %
        'G': 15,  # % Avance Real (INPUT)
        'H': 17,  # PV a corte (INPUT)
        'I': 17,  # AC incurrido (INPUT)
        'J': 17,  # EV ganado (FÓRMULA)
        'K': 17,  # CV (FÓRMULA)
        'L': 17,  # SV (FÓRMULA)
        'M': 13,  # CPI (FÓRMULA)
        'N': 13,  # SPI (FÓRMULA)
        'O': 24,  # Diagnóstico (FÓRMULA)
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Alturas de fila
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 28 # Header row
    for r in range(6, 31):
        ws.row_dimensions[r].height = 22 # Data rows
    ws.row_dimensions[31].height = 10    # Spacing
    ws.row_dimensions[32].height = 28    # Totals row

    ws.views.sheetView[0].showGridLines = True

    # Banner
    ws.merge_cells('B2:O3')
    top_cell = ws['B2']
    title_text = "DATALARIA | CONTROL WBS & SEGUIMIENTO FÍSICO DE ENTREGABLES" if lang == 'ES' else "DATALARIA | WBS CONTROL & PHYSICAL DELIVERABLES TRACKING"
    sub_text = "Estructura de Desglose del Trabajo (25 Paquetes de Trabajo) · Imputación de Avance Real, Varianzas & Eficiencia" if lang == 'ES' else "Work Breakdown Structure (25 Work Packages) · Earned Progress, Variances & Performance Tracking"
    top_cell.value = f"{title_text}\n{sub_text}"
    top_cell.font = font_title
    top_cell.fill = fill_dark
    top_cell.alignment = align_center

    ws['B4'] = "CELDAS BLANCAS = Entradas del Usuario (Avance %, PV, AC) | CELDAS GRISES = Fórmulas EVM Protegidas" if lang == 'ES' else "WHITE CELLS = User Inputs (% Progress, PV, AC) | GRAY CELLS = Protected EVM Formulas"
    ws['B4'].font = Font(name=FONT_NAME, size=9, italic=True, color="64748B")

    # Encabezados
    headers = [
        ("B", "WBS"),
        ("C", "FASE DEL PROYECTO" if lang == 'ES' else "PROJECT PHASE"),
        ("D", "PAQUETE DE TRABAJO (ENTREGABLE)" if lang == 'ES' else "WORK PACKAGE (DELIVERABLE)"),
        ("E", "PRESUPUESTO (BAC)" if lang == 'ES' else "BUDGET (BAC)"),
        ("F", "PESO %" if lang == 'ES' else "WEIGHT %"),
        ("G", "% AVANCE REAL" if lang == 'ES' else "% ACTUAL PROG."),
        ("H", "VALOR PLAN. (PV)" if lang == 'ES' else "PLANNED VAL. (PV)"),
        ("I", "COSTE REAL (AC)" if lang == 'ES' else "ACTUAL COST (AC)"),
        ("J", "VALOR GANADO (EV)" if lang == 'ES' else "EARNED VAL. (EV)"),
        ("K", "VAR. COSTE (CV)" if lang == 'ES' else "COST VAR. (CV)"),
        ("L", "VAR. PLAZO (SV)" if lang == 'ES' else "SCHED. VAR. (SV)"),
        ("M", "CPI" if lang == 'ES' else "CPI"),
        ("N", "SPI" if lang == 'ES' else "SPI"),
        ("O", "DIAGNÓSTICO" if lang == 'ES' else "DIAGNOSIS"),
    ]

    for col_letter, h_text in headers:
        cell = ws[f'{col_letter}5']
        cell.value = h_text
        cell.font = font_tbl_header
        cell.fill = fill_dark
        cell.alignment = align_center
        cell.border = border_cell

    phases = get_wbs_data(lang)
    current_row = 6
    input_cells = []

    for phase_name, tasks in phases:
        for wbs_code, task_name, bac_val, prog_val, pv_val, ac_val in tasks:
            r = current_row

            ws[f'B{r}'] = wbs_code
            ws[f'B{r}'].alignment = align_center
            ws[f'B{r}'].font = font_body_bold
            ws[f'B{r}'].fill = fill_formula
            ws[f'B{r}'].border = border_cell

            ws[f'C{r}'] = phase_name
            ws[f'C{r}'].alignment = align_left
            ws[f'C{r}'].font = font_body
            ws[f'C{r}'].fill = fill_formula
            ws[f'C{r}'].border = border_cell

            ws[f'D{r}'] = task_name
            ws[f'D{r}'].alignment = align_left
            ws[f'D{r}'].font = font_body_bold
            ws[f'D{r}'].fill = fill_formula
            ws[f'D{r}'].border = border_cell

            # BAC (Input editable)
            ws[f'E{r}'] = bac_val
            ws[f'E{r}'].alignment = align_right
            ws[f'E{r}'].font = font_body_bold
            ws[f'E{r}'].number_format = fmt_curr
            ws[f'E{r}'].fill = fill_input
            ws[f'E{r}'].border = border_cell
            input_cells.append(f'E{r}')

            # Peso Ponderado %
            ws[f'F{r}'] = f'=E{r}/$E$32'
            ws[f'F{r}'].alignment = align_right
            ws[f'F{r}'].font = font_body
            ws[f'F{r}'].number_format = fmt_pct
            ws[f'F{r}'].fill = fill_formula
            ws[f'F{r}'].border = border_cell

            # % Avance Real Físico (Input editable)
            ws[f'G{r}'] = prog_val
            ws[f'G{r}'].alignment = align_right
            ws[f'G{r}'].font = font_input
            ws[f'G{r}'].number_format = fmt_pct
            ws[f'G{r}'].fill = fill_input
            ws[f'G{r}'].border = border_cell
            input_cells.append(f'G{r}')

            # PV a corte (Input editable)
            ws[f'H{r}'] = pv_val
            ws[f'H{r}'].alignment = align_right
            ws[f'H{r}'].font = font_body
            ws[f'H{r}'].number_format = fmt_curr
            ws[f'H{r}'].fill = fill_input
            ws[f'H{r}'].border = border_cell
            input_cells.append(f'H{r}')

            # AC incurrido (Input editable)
            ws[f'I{r}'] = ac_val
            ws[f'I{r}'].alignment = align_right
            ws[f'I{r}'].font = font_body
            ws[f'I{r}'].number_format = fmt_curr
            ws[f'I{r}'].fill = fill_input
            ws[f'I{r}'].border = border_cell
            input_cells.append(f'I{r}')

            # EV Ganado: Formula = BAC * % Avance
            ws[f'J{r}'] = f'=E{r}*G{r}'
            ws[f'J{r}'].alignment = align_right
            ws[f'J{r}'].font = font_body_bold
            ws[f'J{r}'].number_format = fmt_curr
            ws[f'J{r}'].fill = fill_formula
            ws[f'J{r}'].border = border_cell

            # CV: Formula = EV - AC
            ws[f'K{r}'] = f'=J{r}-I{r}'
            ws[f'K{r}'].alignment = align_right
            ws[f'K{r}'].font = font_body_bold
            ws[f'K{r}'].number_format = fmt_curr
            ws[f'K{r}'].fill = fill_formula
            ws[f'K{r}'].border = border_cell

            # SV: Formula = EV - PV
            ws[f'L{r}'] = f'=J{r}-H{r}'
            ws[f'L{r}'].alignment = align_right
            ws[f'L{r}'].font = font_body_bold
            ws[f'L{r}'].number_format = fmt_curr
            ws[f'L{r}'].fill = fill_formula
            ws[f'L{r}'].border = border_cell

            # CPI: Formula = IF(AC>0, EV/AC, 1) (usando coma ,)
            ws[f'M{r}'] = f'=IF(I{r}>0, J{r}/I{r}, 1)'
            ws[f'M{r}'].alignment = align_right
            ws[f'M{r}'].font = font_body_bold
            ws[f'M{r}'].number_format = fmt_idx
            ws[f'M{r}'].fill = fill_formula
            ws[f'M{r}'].border = border_cell

            # SPI: Formula = IF(PV>0, EV/PV, 1) (usando coma ,)
            ws[f'N{r}'] = f'=IF(H{r}>0, J{r}/H{r}, 1)'
            ws[f'N{r}'].alignment = align_right
            ws[f'N{r}'].font = font_body_bold
            ws[f'N{r}'].number_format = fmt_idx
            ws[f'N{r}'].fill = fill_formula
            ws[f'N{r}'].border = border_cell

            # Diagnóstico (usando coma ,)
            diag_f = (
                f'=IF(AND(M{r}>=1,N{r}>=1),"En Tiempo & Presupuesto",'
                f'IF(AND(M{r}<0.95,N{r}<0.95),"Crítico (Sobrecoste & Retraso)",'
                f'IF(M{r}<0.95,"Alerta Sobrecoste",'
                f'IF(N{r}<0.95,"Alerta Retraso","En Control"))))'
                if lang == 'ES' else
                f'=IF(AND(M{r}>=1,N{r}>=1),"On Time & Budget",'
                f'IF(AND(M{r}<0.95,N{r}<0.95),"Critical (Overrun & Delay)",'
                f'IF(M{r}<0.95,"Cost Overrun Alert",'
                f'IF(N{r}<0.95,"Schedule Delay Alert","In Control"))))'
            )
            ws[f'O{r}'] = diag_f
            ws[f'O{r}'].alignment = align_center
            ws[f'O{r}'].font = font_body
            ws[f'O{r}'].fill = fill_formula
            ws[f'O{r}'].border = border_cell

            current_row += 1

    # Fila de Totales en Fila 32
    total_row = 32
    ws.merge_cells(f'B{total_row}:D{total_row}')
    ws[f'B{total_row}'] = "TOTAL CONSOLIDADO DEL PROYECTO" if lang == 'ES' else "CONSOLIDATED PROJECT TOTALS"
    ws[f'B{total_row}'].font = font_tbl_header
    ws[f'B{total_row}'].fill = fill_dark
    ws[f'B{total_row}'].alignment = align_center

    # SUM BAC
    ws[f'E{total_row}'] = f'=SUM(E6:E30)'
    ws[f'E{total_row}'].font = font_tbl_header
    ws[f'E{total_row}'].fill = fill_dark
    ws[f'E{total_row}'].number_format = fmt_curr
    ws[f'E{total_row}'].alignment = align_right

    # SUM % (100%)
    ws[f'F{total_row}'] = f'=SUM(F6:F30)'
    ws[f'F{total_row}'].font = font_tbl_header
    ws[f'F{total_row}'].fill = fill_dark
    ws[f'F{total_row}'].number_format = fmt_pct
    ws[f'F{total_row}'].alignment = align_right

    # % Avance Físico Global = EV / BAC
    ws[f'G{total_row}'] = f'=J{total_row}/E{total_row}'
    ws[f'G{total_row}'].font = font_tbl_header
    ws[f'G{total_row}'].fill = fill_dark
    ws[f'G{total_row}'].number_format = fmt_pct
    ws[f'G{total_row}'].alignment = align_right

    # SUM PV
    ws[f'H{total_row}'] = f'=SUM(H6:H30)'
    ws[f'H{total_row}'].font = font_tbl_header
    ws[f'H{total_row}'].fill = fill_dark
    ws[f'H{total_row}'].number_format = fmt_curr
    ws[f'H{total_row}'].alignment = align_right

    # SUM AC
    ws[f'I{total_row}'] = f'=SUM(I6:I30)'
    ws[f'I{total_row}'].font = font_tbl_header
    ws[f'I{total_row}'].fill = fill_dark
    ws[f'I{total_row}'].number_format = fmt_curr
    ws[f'I{total_row}'].alignment = align_right

    # SUM EV
    ws[f'J{total_row}'] = f'=SUM(J6:J30)'
    ws[f'J{total_row}'].font = font_tbl_header
    ws[f'J{total_row}'].fill = fill_dark
    ws[f'J{total_row}'].number_format = fmt_curr
    ws[f'J{total_row}'].alignment = align_right

    # SUM CV
    ws[f'K{total_row}'] = f'=J{total_row}-I{total_row}'
    ws[f'K{total_row}'].font = font_tbl_header
    ws[f'K{total_row}'].fill = fill_dark
    ws[f'K{total_row}'].number_format = fmt_curr
    ws[f'K{total_row}'].alignment = align_right

    # SUM SV
    ws[f'L{total_row}'] = f'=J{total_row}-H{total_row}'
    ws[f'L{total_row}'].font = font_tbl_header
    ws[f'L{total_row}'].fill = fill_dark
    ws[f'L{total_row}'].number_format = fmt_curr
    ws[f'L{total_row}'].alignment = align_right

    # GLOBAL CPI
    ws[f'M{total_row}'] = f'=IF(I{total_row}>0, J{total_row}/I{total_row}, 1)'
    ws[f'M{total_row}'].font = font_tbl_header
    ws[f'M{total_row}'].fill = fill_dark
    ws[f'M{total_row}'].number_format = fmt_idx
    ws[f'M{total_row}'].alignment = align_right

    # GLOBAL SPI
    ws[f'N{total_row}'] = f'=IF(H{total_row}>0, J{total_row}/H{total_row}, 1)'
    ws[f'N{total_row}'].font = font_tbl_header
    ws[f'N{total_row}'].fill = fill_dark
    ws[f'N{total_row}'].number_format = fmt_idx
    ws[f'N{total_row}'].alignment = align_right

    # GLOBAL DIAGNOSIS
    ws[f'O{total_row}'] = f'=IF(AND(M{total_row}>=1,N{total_row}>=1),"PROYECTO EN SALUD","DESVIACIÓN CRÍTICA")' if lang == 'ES' else f'=IF(AND(M{total_row}>=1,N{total_row}>=1),"HEALTHY PROJECT","CRITICAL VARIANCE")'
    ws[f'O{total_row}'].font = font_tbl_header
    ws[f'O{total_row}'].fill = fill_dark
    ws[f'O{total_row}'].alignment = align_center

    for col_idx in range(2, 16):
        col_letter = get_column_letter(col_idx)
        ws[f'{col_letter}{total_row}'].border = border_top_total

    return input_cells


def build_tab3_timeseries(ws, lang='ES'):
    """
    Construye la Pestaña 3: Serie Temporal & Proyecciones EAC.
    4 Columnas contiguas para el gráfico de Curva S (Cols F, G, H, I):
    - Col F: Σ PV ACUM.
    - Col G: Σ EV ACUM.
    - Col H: Σ AC ACUM.
    - Col I: PROYECCIÓN EAC
    - Col J: CV ACUM.
    - Col K: SV ACUM.
    - Col L: CPI ACUM.
    - Col M: SPI ACUM.
    Todas las fórmulas usan comas ',' (estándar OpenXML ECMA-376).
    """
    fmt_curr = '#,##0 "€"' if lang == 'ES' else '"$"#,##0'
    fmt_idx = '0.00"x"'
    fmt_pct = '0.0%'

    col_widths = {
        'A': 3,
        'B': 14,  # Periodo / Parámetro
        'C': 22,  # PV Mensual / Fórmula Fuente
        'D': 18,  # EV Mensual / Valor Parámetro
        'E': 16,  # AC Mensual
        'F': 18,  # PV Acumulado (Chart S1)
        'G': 18,  # EV Acumulado (Chart S2)
        'H': 18,  # AC Acumulado (Chart S3)
        'I': 20,  # Proyección EAC (Chart S4)
        'J': 16,  # CV Acumulado
        'K': 16,  # SV Acumulado
        'L': 14,  # CPI Acumulado
        'M': 14,  # SPI Acumulado
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Alturas de fila explícitas
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 28 # Header row
    for r in range(6, 18):
        ws.row_dimensions[r].height = 22 # Data rows
    ws.row_dimensions[18].height = 10
    ws.row_dimensions[19].height = 10
    ws.row_dimensions[20].height = 26 # Section header row
    for r in range(21, 30):
        ws.row_dimensions[r].height = 24 # Parameters rows

    ws.views.sheetView[0].showGridLines = True

    # Banner
    ws.merge_cells('B2:M3')
    top_cell = ws['B2']
    title_text = "DATALARIA | EVOLUCIÓN TEMPORAL, CURVA S & MODELOS DE PROYECCIÓN EAC" if lang == 'ES' else "DATALARIA | TIME SERIES, S-CURVE & MATHEMATICAL EAC FORECASTING"
    sub_text = "Registro Mensual (Mes 1 a Mes 12) · Modelos Típico, Atípico y Compuesto · Índice de Rendimiento TCPI" if lang == 'ES' else "Monthly Tracking (M01 to M12) · Typical, Atypical & Combined EAC Models · TCPI Feasibility Index"
    top_cell.value = f"{title_text}\n{sub_text}"
    top_cell.font = font_title
    top_cell.fill = fill_dark
    top_cell.alignment = align_center

    # Encabezados
    headers = [
        ("B", "MES" if lang == 'ES' else "MONTH"),
        ("C", "PV MENSUAL" if lang == 'ES' else "MONTHLY PV"),
        ("D", "EV MENSUAL" if lang == 'ES' else "MONTHLY EV"),
        ("E", "AC MENSUAL" if lang == 'ES' else "MONTHLY AC"),
        ("F", "Σ PV ACUM." if lang == 'ES' else "CUMUL. PV"),
        ("G", "Σ EV ACUM." if lang == 'ES' else "CUMUL. EV"),
        ("H", "Σ AC ACUM." if lang == 'ES' else "CUMUL. AC"),
        ("I", "PROYECCIÓN EAC" if lang == 'ES' else "EAC FORECAST"),
        ("J", "CV ACUM." if lang == 'ES' else "CUMUL. CV"),
        ("K", "SV ACUM." if lang == 'ES' else "CUMUL. SV"),
        ("L", "CPI ACUM." if lang == 'ES' else "CUMUL. CPI"),
        ("M", "SPI ACUM." if lang == 'ES' else "CUMUL. SPI"),
    ]

    for col_letter, h_text in headers:
        cell = ws[f'{col_letter}5']
        cell.value = h_text
        cell.font = font_tbl_header
        cell.fill = fill_dark
        cell.alignment = align_center
        cell.border = border_cell

    monthly_records = get_monthly_data(lang)
    input_cells = []

    for idx, (m_code, pv_m, ev_m, ac_m) in enumerate(monthly_records):
        r = 6 + idx

        ws[f'B{r}'] = m_code
        ws[f'B{r}'].alignment = align_center
        ws[f'B{r}'].font = font_body_bold
        ws[f'B{r}'].fill = fill_formula
        ws[f'B{r}'].border = border_cell

        # PV Mensual (Input)
        ws[f'C{r}'] = pv_m
        ws[f'C{r}'].alignment = align_right
        ws[f'C{r}'].font = font_body
        ws[f'C{r}'].number_format = fmt_curr
        ws[f'C{r}'].fill = fill_input
        ws[f'C{r}'].border = border_cell
        input_cells.append(f'C{r}')

        # EV Mensual (Input para M1-M6)
        ws[f'D{r}'] = ev_m if ev_m != "" else ""
        ws[f'D{r}'].alignment = align_right
        ws[f'D{r}'].font = font_input if ev_m != "" else font_body
        if ev_m != "":
            ws[f'D{r}'].number_format = fmt_curr
            ws[f'D{r}'].fill = fill_input
            input_cells.append(f'D{r}')
        else:
            ws[f'D{r}'].fill = fill_formula
        ws[f'D{r}'].border = border_cell

        # AC Mensual (Input para M1-M6)
        ws[f'E{r}'] = ac_m if ac_m != "" else ""
        ws[f'E{r}'].alignment = align_right
        ws[f'E{r}'].font = font_body
        if ac_m != "":
            ws[f'E{r}'].number_format = fmt_curr
            ws[f'E{r}'].fill = fill_input
            input_cells.append(f'E{r}')
        else:
            ws[f'E{r}'].fill = fill_formula
        ws[f'E{r}'].border = border_cell

        # PV Acumulado = SUM(C$6:C{r})
        ws[f'F{r}'] = f'=SUM(C$6:C{r})'
        ws[f'F{r}'].alignment = align_right
        ws[f'F{r}'].font = font_body
        ws[f'F{r}'].number_format = fmt_curr
        ws[f'F{r}'].fill = fill_formula
        ws[f'F{r}'].border = border_cell

        # EV Acumulado: en M1-M6 calcula la suma; en M7-M12 devuelve NA() para que no trace en gráfica
        if idx <= 5:
            ws[f'G{r}'] = f'=SUM(D$6:D{r})'
            ws[f'G{r}'].number_format = fmt_curr
        else:
            ws[f'G{r}'] = '=NA()'
        ws[f'G{r}'].alignment = align_right
        ws[f'G{r}'].font = font_body_bold
        ws[f'G{r}'].fill = fill_formula
        ws[f'G{r}'].border = border_cell

        # AC Acumulado: en M1-M6 calcula la suma; en M7-M12 devuelve NA()
        if idx <= 5:
            ws[f'H{r}'] = f'=SUM(E$6:E{r})'
            ws[f'H{r}'].number_format = fmt_curr
        else:
            ws[f'H{r}'] = '=NA()'
        ws[f'H{r}'].alignment = align_right
        ws[f'H{r}'].font = font_body_bold
        ws[f'H{r}'].fill = fill_formula
        ws[f'H{r}'].border = border_cell

        # Proyección EAC (Col I): M1-M5 es NA(); M6 arranca en H11; M7-M12 interpola hasta D23 (EAC 1)
        if idx < 5:
            ws[f'I{r}'] = '=NA()'
        elif idx == 5:
            ws[f'I{r}'] = f'=H{r}'
            ws[f'I{r}'].number_format = fmt_curr
        else:
            step = idx - 5
            ws[f'I{r}'] = f'=H$11 + ($D$23 - H$11) * ({step}/6)'
            ws[f'I{r}'].number_format = fmt_curr
        ws[f'I{r}'].alignment = align_right
        ws[f'I{r}'].font = font_body_bold
        ws[f'I{r}'].fill = fill_formula
        ws[f'I{r}'].border = border_cell

        # CV Acumulado (Col J): en M1-M6 = G{r}-H{r}; M7-M12 vacío
        if idx <= 5:
            ws[f'J{r}'] = f'=G{r}-H{r}'
            ws[f'J{r}'].number_format = fmt_curr
        else:
            ws[f'J{r}'] = '=""'
        ws[f'J{r}'].alignment = align_right
        ws[f'J{r}'].font = font_body
        ws[f'J{r}'].fill = fill_formula
        ws[f'J{r}'].border = border_cell

        # SV Acumulado (Col K): en M1-M6 = G{r}-F{r}; M7-M12 vacío
        if idx <= 5:
            ws[f'K{r}'] = f'=G{r}-F{r}'
            ws[f'K{r}'].number_format = fmt_curr
        else:
            ws[f'K{r}'] = '=""'
        ws[f'K{r}'].alignment = align_right
        ws[f'K{r}'].font = font_body
        ws[f'K{r}'].fill = fill_formula
        ws[f'K{r}'].border = border_cell

        # CPI Acumulado (Col L): en M1-M6 = G{r}/H{r} con IF(H>0); M7-M12 vacío
        if idx <= 5:
            ws[f'L{r}'] = f'=IF(H{r}>0, G{r}/H{r}, 1)'
            ws[f'L{r}'].number_format = fmt_idx
        else:
            ws[f'L{r}'] = '=""'
        ws[f'L{r}'].alignment = align_right
        ws[f'L{r}'].font = font_body_bold
        ws[f'L{r}'].fill = fill_formula
        ws[f'L{r}'].border = border_cell

        # SPI Acumulado (Col M): en M1-M6 = G{r}/F{r} con IF(F>0); M7-M12 vacío
        if idx <= 5:
            ws[f'M{r}'] = f'=IF(F{r}>0, G{r}/F{r}, 1)'
            ws[f'M{r}'].number_format = fmt_idx
        else:
            ws[f'M{r}'] = '=""'
        ws[f'M{r}'].alignment = align_right
        ws[f'M{r}'].font = font_body_bold
        ws[f'M{r}'].fill = fill_formula
        ws[f'M{r}'].border = border_cell

    # 2. MODELOS MATEMÁTICOS DE ESTIMACIÓN A LA FINALIZACIÓN (EAC)
    # Filas 20 a 29
    ws.merge_cells('B20:D20')
    ws['B20'] = "MODELOS COMPARATIVOS DE ESTIMACIÓN EAC & TCPI" if lang == 'ES' else "COMPARATIVE EAC FORECASTING MODELS & TCPI"
    ws['B20'].font = font_section
    ws['B20'].fill = fill_dark
    ws['B20'].alignment = align_left

    ref_wbs = "'Control WBS & Entregables'" if lang == 'ES' else "'WBS & Deliverables Tracking'"

    # Tabla de parámetros base: Col B (Parámetro), Col C (Fórmula descrita), Col D (Fórmula calculada)
    # L11 es CPI en M06, M11 es SPI en M06, G11 es EV en M06, H11 es AC en M06, F11 es PV en M06.
    models_def = [
        ("B21", "C21", "D21", "PARÁMETRO" if lang == 'ES' else "PARAMETER", "FÓRMULA / FUENTE" if lang == 'ES' else "FORMULA / SOURCE", "VALOR" if lang == 'ES' else "VALUE"),
        ("B22", "C22", "D22", "Presupuesto Total (BAC)" if lang == 'ES' else "Budget at Completion (BAC)", "Línea Base Aprobada (WBS)" if lang == 'ES' else "Approved Baseline (WBS)", f"={ref_wbs}!E32"),
        ("B23", "C23", "D23", "EAC 1 (CPI Típico)" if lang == 'ES' else "EAC 1 (Typical CPI)", "BAC / CPI", "=D22/L11"),
        ("B24", "C24", "D24", "EAC 2 (Atípico)" if lang == 'ES' else "EAC 2 (Atypical Rest)", "AC + (BAC - EV)", "=H11 + (D22 - G11)"),
        ("B25", "C25", "D25", "EAC 3 (Compuesto)" if lang == 'ES' else "EAC 3 (Combined CPIxSPI)", "AC + (BAC - EV)/(CPI*SPI)", "=H11 + (D22 - G11)/(L11*M11)"),
        ("B26", "C26", "D26", "Varianza Cierre VAC 1" if lang == 'ES' else "Variance at Close VAC 1", "BAC - EAC 1", "=D22 - D23"),
        ("B27", "C27", "D27", "Varianza Cierre VAC 3" if lang == 'ES' else "Variance at Close VAC 3", "BAC - EAC 3", "=D22 - D25"),
        ("B28", "C28", "D28", "TCPI (BAC Original)" if lang == 'ES' else "TCPI (To Achieve BAC)", "(BAC - EV) / (BAC - AC)", "=(D22 - G11)/(D22 - H11)"),
        ("B29", "C29", "D29", "TCPI (EAC 1 Revisado)" if lang == 'ES' else "TCPI (To Achieve EAC 1)", "(BAC - EV) / (EAC 1 - AC)", "=(D22 - G11)/(D23 - H11)"),
    ]

    for c_p, c_f, c_v, param, formula_desc, val_formula in models_def:
        ws[c_p] = param
        ws[c_p].font = font_tbl_subheader if "PARÁMETRO" in param or "PARAMETER" in param else font_body_bold
        ws[c_p].fill = fill_dark if "PARÁMETRO" in param or "PARAMETER" in param else fill_card
        ws[c_p].border = border_cell

        ws[c_f] = formula_desc
        ws[c_f].font = font_tbl_subheader if "PARÁMETRO" in param or "PARAMETER" in param else font_body
        ws[c_f].fill = fill_dark if "PARÁMETRO" in param or "PARAMETER" in param else fill_card
        ws[c_f].border = border_cell

        ws[c_v] = val_formula
        ws[c_v].font = font_tbl_subheader if "PARÁMETRO" in param or "PARAMETER" in param else font_body_bold
        ws[c_v].fill = fill_dark if "PARÁMETRO" in param or "PARAMETER" in param else fill_card
        ws[c_v].border = border_cell

        if "EAC" in param or "BAC" in param or "Varianza" in param or "Variance" in param:
            ws[c_v].number_format = fmt_curr
        elif "TCPI" in param:
            ws[c_v].number_format = fmt_idx

    # Panel lateral de gobernanza TCPI (Cols E a M en filas 20 a 29)
    ws.merge_cells('E20:M29')
    tcpi_box = ws['E20']
    tcpi_text = (
        "ANÁLISIS DE VIABILIDAD DIRECTIVA TCPI (PMBOK / ANSI-748):\n\n"
        "1. TCPI (BAC) = 1.14x:\n"
        "   Para recuperar el presupuesto original de 1.200.000 €, el equipo de proyecto debería alcanzar\n"
        "   una eficiencia de ejecución del 114% en todo el trabajo restante (un salto de productividad del 32,5%).\n"
        "   -> CONCLUSIÓN: Matemáticamente inviable según el histórico empírico del PMBOK (umbral crítico: >1.10x).\n\n"
        "2. TCPI (EAC 1) = 0.86x:\n"
        "   Si el Comité aprueba la reprogramación presupuestaria al techo EAC de 1.395.349 €,\n"
        "   el proyecto es 100% viable manteniendo la tasa de productividad observada.\n\n"
        "3. RESOLUCIÓN REQUERIDA:\n"
        "   Aprobar un re-baselining presupuestario con inyección de 195.349 € de reservas de gestión\n"
        "   combinado con Fast-Tracking en Fase 3 para amortiguar el impacto del retraso en el camino crítico."
        if lang == 'ES' else
        "TCPI GOVERNANCE & FEASIBILITY ANALYSIS (PMBOK / ANSI-748):\n\n"
        "1. TCPI (BAC) = 1.14x:\n"
        "   To recover the original $1,200,000 budget, the project team would need to operate at 114% efficiency\n"
        "   across all remaining deliverables (a 32.5% productivity surge over current performance).\n"
        "   -> CONCLUSION: Statistically unviable per PMBOK empirical benchmarks (critical cutoff: >1.10x).\n\n"
        "2. TCPI (EAC 1) = 0.86x:\n"
        "   If the Executive Board ratifies a revised budget baseline at EAC 1 ($1,395,349),\n"
        "   the remaining project scope is 100% achievable at the current demonstrated burn rate.\n\n"
        "3. REQUIRED RESOLUTION:\n"
        "   Formally re-baseline the budget with a $195,349 management reserve release,\n"
        "   paired with Fast-Tracking in Phase 3 to curb critical path schedule delay."
    )
    tcpi_box.value = tcpi_text
    tcpi_box.font = Font(name=FONT_NAME, size=9, bold=False, color=COLOR_NAVY_DARK)
    tcpi_box.fill = fill_blue_card
    tcpi_box.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for r in range(20, 30):
        for c in range(5, 14):
            ws.cell(row=r, column=c).border = border_cell

    return input_cells


def build_tab4_action_plan(ws, lang='ES'):
    """Construye la Pestaña 4: Plan de Acción & Recuperación."""
    col_widths = {
        'A': 3,
        'B': 8,   # ID
        'C': 34,  # Entregable
        'D': 25,  # Diagnóstico
        'E': 22,  # Palanca
        'F': 44,  # Medida Detallada
        'G': 18,  # Responsable
        'H': 14,  # Delta CPI
        'I': 14,  # Delta SPI
        'J': 14,  # Fecha Límite
        'K': 18,  # Estado
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 28 # Header row
    for r in range(6, 15):
        ws.row_dimensions[r].height = 26 # Action rows with ample breathing room
    ws.row_dimensions[15].height = 28    # Totals row

    ws.views.sheetView[0].showGridLines = True

    # Banner
    ws.merge_cells('B2:K3')
    top_cell = ws['B2']
    title_text = "DATALARIA | PLAN DE ACCIÓN, MITIGACIÓN & RECUPERACIÓN OPERATIVA" if lang == 'ES' else "DATALARIA | ACTION PLAN, MITIGATION & OPERATIONAL RECOVERY"
    sub_text = "Matriz de Intervención sobre Paquetes de Trabajo en Crisis · Crashing, Fast-Tracking, Descope & Compromisos C-Level" if lang == 'ES' else "Intervention Matrix for At-Risk Deliverables · Crashing, Fast-Tracking, Descope & C-Level Sign-Offs"
    top_cell.value = f"{title_text}\n{sub_text}"
    top_cell.font = font_title
    top_cell.fill = fill_dark
    top_cell.alignment = align_center

    headers = [
        ("B", "ID"),
        ("C", "ENTREGABLE EN DESVIACIÓN" if lang == 'ES' else "AT-RISK DELIVERABLE"),
        ("D", "DIAGNÓSTICO DESVÍO" if lang == 'ES' else "ROOT CAUSE / VARIANCE"),
        ("E", "PALANCA DE AJUSTE" if lang == 'ES' else "INTERVENTION LEVER"),
        ("F", "MEDIDA OPERATIVA DETALLADA" if lang == 'ES' else "OPERATIONAL ACTION PLAN"),
        ("G", "RESPONSABLE C-LEVEL" if lang == 'ES' else "C-LEVEL OWNER"),
        ("H", "Δ CPI ESPERADO" if lang == 'ES' else "EXP. Δ CPI"),
        ("I", "Δ SPI ESPERADO" if lang == 'ES' else "EXP. Δ SPI"),
        ("J", "FECHA LÍMITE" if lang == 'ES' else "DEADLINE"),
        ("K", "ESTADO" if lang == 'ES' else "STATUS"),
    ]

    for col_letter, h_text in headers:
        cell = ws[f'{col_letter}5']
        cell.value = h_text
        cell.font = font_tbl_header
        cell.fill = fill_dark
        cell.alignment = align_center
        cell.border = border_cell

    if lang == 'ES':
        actions = [
            ("ACT-01", "2.1 Motor Transaccional Conciliación", "Sobrecoste +20k€ y complejidad técnica", "Crashing Selectivo", "Contratación de 2 arquitectos senior especialistas por 6 semanas", "CTO / Lead Arch", "+0.04", "+0.05", "15/Nov/2026", "Aprobado Board"),
            ("ACT-02", "2.3 Pasarelas de Pago & Routing", "Retraso en homologación bancaria externa", "Fast-Tracking", "Paralelizar pruebas de sandbox mientras se audita la seguridad", "COO / PMO Head", "+0.02", "+0.06", "30/Nov/2026", "En Ejecución"),
            ("ACT-03", "3.1 Conector ERP Core SAP", "Desviación de alcance y tarifas de consultor", "Renegociación Tarifas", "Fijar precio cerrado para deliverables restantes y limitar T&M", "CFO / Procurement", "+0.05", "+0.02", "20/Nov/2026", "Negociación"),
            ("ACT-04", "3.3 Conector CRM Salesforce", "Bajo valor crítico para el Go-Live inicial", "Descope Negociado", "Diferir fase 2 de sincronización bidireccional a Q1 post-lanzamiento", "Chief Commercial Off.", "+0.03", "+0.04", "10/Nov/2026", "Aprobado Board"),
            ("ACT-05", "3.2 Pasarelas Bancarias & SWIFT", "Dependencia de validación con banco corresponsal", "Escalado Ejecutivo", "Reunión bilateral C-Level con Director General de la entidad bancaria", "CEO / Sponsor", "+0.01", "+0.04", "05/Nov/2026", "Agendado"),
            ("ACT-06", "4.1 Migración de Datos Históricos", "Volumetría no optimizada en esquemas legacy", "Automatización Scripts", "Reutilizar framework de ingesta paralelizado con Apache Spark", "Data Eng Director", "+0.02", "+0.03", "15/Dic/2026", "En Curso"),
            ("ACT-07", "4.3 Pruebas de Carga & Estrés", "Riesgo de cuello de botella en entorno UAT", "IaC Escalado Cloud", "Aprovisionamiento spot instances con escalado dinámico programado", "Head of DevOps", "+0.01", "+0.02", "20/Dic/2026", "Planificado"),
            ("ACT-08", "Gobernanza & Reporte Quincenal", "Falta de visibilidad temprana de varianzas", "Ritmo de Control EVM", "Comité de control de valor ganado quincenal obligatorio con el CFO", "CFO / COO", "+0.03", "+0.03", "Inmediato", "Implantado"),
        ]
    else:
        actions = [
            ("ACT-01", "2.1 Core Reconciliation Engine", "+$20k cost overrun & complex algorithms", "Selective Crashing", "Engage 2 senior distributed systems architects for 6 weeks", "CTO / Lead Arch", "+0.04", "+0.05", "15/Nov/2026", "Board Approved"),
            ("ACT-02", "2.3 Payment Gateway & Routing", "External bank compliance delays", "Fast-Tracking", "Execute parallel sandbox testing alongside compliance sign-off", "COO / PMO Head", "+0.02", "+0.06", "30/Nov/2026", "In Progress"),
            ("ACT-03", "3.1 Core ERP Connector (SAP)", "Scope creep and contractor rate inflation", "Vendor Renegotiation", "Transition remaining milestones to fixed-fee cap contracts", "CFO / Procurement", "+0.05", "+0.02", "20/Nov/2026", "Negotiation"),
            ("ACT-04", "3.3 CRM Salesforce Integration", "Non-critical path for initial Go-Live", "Negotiated Descope", "Postpone bidirectional contact sync to Phase 2 (post Go-Live)", "Chief Commercial Off.", "+0.03", "+0.04", "10/Nov/2026", "Board Approved"),
            ("ACT-05", "3.2 Banking Gateway & SWIFT", "Correspondent banking certification blocked", "Executive Escalation", "Bilateral C-Level alignment with Partner Bank Managing Director", "CEO / Sponsor", "+0.01", "+0.04", "05/Nov/2026", "Scheduled"),
            ("ACT-06", "4.1 Legacy Data Migration", "Legacy data bloat and unindexed schemas", "Pipeline Automation", "Deploy parallel PySpark extraction pipelines to reduce runtimes", "Data Eng Director", "+0.02", "+0.03", "15/Dec/2026", "In Progress"),
            ("ACT-07", "4.3 Performance & Load Testing", "Risk of infrastructure bottlenecks", "Cloud Spot Scaling", "Provision elastic spot nodes with automated cluster spin-down", "Head of DevOps", "+0.01", "+0.02", "20/Dec/2026", "Planned"),
            ("ACT-08", "Bi-Weekly EVM Governance", "Late detection of deliverable variance", "EVM Review Cadence", "Mandatory bi-weekly EVM audit cadence directly with the CFO", "CFO / COO", "+0.03", "+0.03", "Immediate", "Implemented"),
        ]

    input_cells = []
    for idx, (aid, deliverable, root_cause, lever, action_desc, owner, dcpi, dspi, dline, status) in enumerate(actions):
        r = 6 + idx

        ws[f'B{r}'] = aid
        ws[f'B{r}'].alignment = align_center
        ws[f'B{r}'].font = font_body_bold
        ws[f'B{r}'].fill = fill_formula
        ws[f'B{r}'].border = border_cell

        ws[f'C{r}'] = deliverable
        ws[f'C{r}'].alignment = align_left
        ws[f'C{r}'].font = font_body_bold
        ws[f'C{r}'].fill = fill_input
        ws[f'C{r}'].border = border_cell
        input_cells.append(f'C{r}')

        ws[f'D{r}'] = root_cause
        ws[f'D{r}'].alignment = align_left
        ws[f'D{r}'].font = font_body
        ws[f'D{r}'].fill = fill_input
        ws[f'D{r}'].border = border_cell
        input_cells.append(f'D{r}')

        ws[f'E{r}'] = lever
        ws[f'E{r}'].alignment = align_center
        ws[f'E{r}'].font = font_body_bold
        ws[f'E{r}'].fill = fill_input
        ws[f'E{r}'].border = border_cell
        input_cells.append(f'E{r}')

        ws[f'F{r}'] = action_desc
        ws[f'F{r}'].alignment = align_left
        ws[f'F{r}'].font = font_body
        ws[f'F{r}'].fill = fill_input
        ws[f'F{r}'].border = border_cell
        input_cells.append(f'F{r}')

        ws[f'G{r}'] = owner
        ws[f'G{r}'].alignment = align_center
        ws[f'G{r}'].font = font_body_bold
        ws[f'G{r}'].fill = fill_input
        ws[f'G{r}'].border = border_cell
        input_cells.append(f'G{r}')

        ws[f'H{r}'] = dcpi
        ws[f'H{r}'].alignment = align_right
        ws[f'H{r}'].font = font_input
        ws[f'H{r}'].fill = fill_input
        ws[f'H{r}'].border = border_cell
        input_cells.append(f'H{r}')

        ws[f'I{r}'] = dspi
        ws[f'I{r}'].alignment = align_right
        ws[f'I{r}'].font = font_input
        ws[f'I{r}'].fill = fill_input
        ws[f'I{r}'].border = border_cell
        input_cells.append(f'I{r}')

        ws[f'J{r}'] = dline
        ws[f'J{r}'].alignment = align_center
        ws[f'J{r}'].font = font_body
        ws[f'J{r}'].fill = fill_input
        ws[f'J{r}'].border = border_cell
        input_cells.append(f'J{r}')

        ws[f'K{r}'] = status
        ws[f'K{r}'].alignment = align_center
        ws[f'K{r}'].font = font_body_bold
        ws[f'K{r}'].fill = fill_green if "Aprobado" in status or "Approved" in status or "Implantado" in status or "Implemented" in status else fill_amber
        ws[f'K{r}'].border = border_cell
        input_cells.append(f'K{r}')

    # Fila resumen
    ws.merge_cells('B15:G15')
    ws['B15'] = "IMPACTO CONSOLIDADO ESTIMADO DEL PLAN DE RECUPERACIÓN" if lang == 'ES' else "CONSOLIDATED ESTIMATED RECOVERY IMPACT"
    ws['B15'].font = font_tbl_header
    ws['B15'].fill = fill_dark
    ws['B15'].alignment = align_center

    ws['H15'] = "+0.21"
    ws['H15'].font = font_tbl_header
    ws['H15'].fill = fill_dark
    ws['H15'].alignment = align_right

    ws['I15'] = "+0.29"
    ws['I15'].font = font_tbl_header
    ws['I15'].fill = fill_dark
    ws['I15'].alignment = align_right

    ws.merge_cells('J15:K15')
    ws['J15'] = "RESTAURA CPI ≥ 1.00 & SPI ≥ 1.00" if lang == 'ES' else "RESTORES CPI ≥ 1.00 & SPI ≥ 1.00"
    ws['J15'].font = font_tbl_header
    ws['J15'].fill = fill_dark
    ws['J15'].alignment = align_center

    for col_idx in range(2, 12):
        col_letter = get_column_letter(col_idx)
        ws[f'{col_letter}15'].border = border_top_total

    return input_cells


def add_scurve_chart_to_dashboard(ws_dash, ws_ts, lang='ES'):
    """
    Agrega el gráfico de la Curva S ejecutiva nativo en la Pestaña 1 (Dashboard).
    Datos contiguos desde la Pestaña 3:
    - Series: Cols F (Σ PV), G (Σ EV), H (Σ AC), I (PROYECCIÓN EAC) [Cols 6 a 9]
    - Categorías: Col B (M01 a M12) [Col 2, Rows 6 a 17]
    Configuración completa de ejes visibles, etiquetas y colores corporativos.
    """
    chart = LineChart()
    chart.title = None  # El título ejecutivo ya se muestra en la cabecera H18 del Dashboard
    chart.y_axis.title = None  # Los valores numéricos del eje Y ya incluyen el símbolo de divisa (€ / $)
    chart.x_axis.title = None  # Los periodos M01 a M12 son autoexplicativos y evitan colisiones de texto
    chart.width = 21.0
    chart.height = 12.8

    # 4 columnas contiguas: F (6), G (7), H (8), I (9)
    data = Reference(ws_ts, min_col=6, min_row=5, max_col=9, max_row=17)
    chart.add_data(data, titles_from_data=True)

    # Inyectar categorías de texto explícitas mediante StrRef y strCache para renderizado nativo garantizado
    tab_name = "Serie Temporal & Proyecciones" if lang == 'ES' else "Time Series & EAC Forecasting"
    months = [f"M{i:02d}" for i in range(1, 13)]
    str_vals = [StrVal(idx=i, v=m) for i, m in enumerate(months)]
    str_data = StrData(pt=str_vals)
    cat_ref_formula = f"'{tab_name}'!$B$6:$B$17"

    for s in chart.series:
        s.cat = AxDataSource(strRef=StrRef(f=cat_ref_formula, strCache=str_data))

    # Estilizar cada serie con colores corporativos Datalaria
    # Series 0: PV Acumulado (Azul #2563EB)
    # Series 1: EV Acumulado (Verde #10B981)
    # Series 2: AC Acumulado (Rojo #DC2626)
    # Series 3: EAC Proyección (Púrpura #7C3AED, discontinuo)
    colors = ["2563EB", "10B981", "DC2626", "7C3AED"]
    for idx, c_hex in enumerate(colors):
        s = chart.series[idx]
        s.graphicalProperties.line.solidFill = c_hex
        s.graphicalProperties.line.width = 25000 if idx < 3 else 22000
        if idx == 3:
            s.graphicalProperties.line.dashStyle = "dash"

    # Configuración de visibilidad de ejes y números (ECMA-376)
    chart.x_axis.delete = False
    chart.x_axis.axPos = "b"  # Posición inferior (Bottom)
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.crosses = "autoZero"

    chart.y_axis.delete = False
    chart.y_axis.axPos = "l"  # Posición izquierda (Left)
    chart.y_axis.tickLblPos = "nextTo"
    chart.y_axis.crosses = "autoZero"
    chart.y_axis.number_format = '#,##0 €' if lang == 'ES' else '$#,##0'
    chart.y_axis.majorGridlines = ChartLines()
    chart.legend.legendPos = "r"

    ws_dash.add_chart(chart, "H19")


def build_workbook(lang='ES'):
    """Construye el libro completo de 4 pestañas y aplica la política de protección."""
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    if lang == 'ES':
        title_tab1 = "Dashboard Ejecutivo EVM"
        title_tab2 = "Control WBS & Entregables"
        title_tab3 = "Serie Temporal & Proyecciones"
        title_tab4 = "Plan de Acción & Recuperación"
    else:
        title_tab1 = "Executive EVM Dashboard"
        title_tab2 = "WBS & Deliverables Tracking"
        title_tab3 = "Time Series & EAC Forecasting"
        title_tab4 = "Action & Recovery Plan"

    ws1 = wb.create_sheet(title=title_tab1)
    ws2 = wb.create_sheet(title=title_tab2)
    ws3 = wb.create_sheet(title=title_tab3)
    ws4 = wb.create_sheet(title=title_tab4)

    wb.remove(default_sheet)

    print(f"[{lang}] Construyendo Pestaña 1: {title_tab1}...")
    build_tab1_dashboard(ws1, lang=lang)

    print(f"[{lang}] Construyendo Pestaña 2: {title_tab2}...")
    tab2_inputs = build_tab2_wbs(ws2, lang=lang)

    print(f"[{lang}] Construyendo Pestaña 3: {title_tab3}...")
    tab3_inputs = build_tab3_timeseries(ws3, lang=lang)

    print(f"[{lang}] Construyendo Pestaña 4: {title_tab4}...")
    tab4_inputs = build_tab4_action_plan(ws4, lang=lang)

    print(f"[{lang}] Incrustando gráfico nativo de Curva S en Dashboard...")
    add_scurve_chart_to_dashboard(ws1, ws3, lang=lang)

    # Aplicar protección ECMA-376
    print(f"[{lang}] Aplicando política de protección ECMA-376 con contraseña oficial...")
    finalize_protection(ws1, input_ranges=[])
    finalize_protection(ws2, input_ranges=tab2_inputs)
    finalize_protection(ws3, input_ranges=tab3_inputs)
    finalize_protection(ws4, input_ranges=tab4_inputs)

    return wb


def generate_all_excel():
    """Genera las dos versiones del modelo Excel en sus rutas oficiales."""
    dir_base = "packages/Suite_03_Control_Operativo/04_Cuadro_EVM_Valor_Ganado"
    dir_es = os.path.join(dir_base, "[ES]_Cuadro_EVM_Valor_Ganado")
    dir_en = os.path.join(dir_base, "[EN]_Earned_Value_Management_EVM")

    os.makedirs(dir_es, exist_ok=True)
    os.makedirs(dir_en, exist_ok=True)

    path_es = os.path.join(dir_es, "Cuadro_Mando_EVM_Datalaria_ES.xlsx")
    path_en = os.path.join(dir_en, "EVM_Dashboard_Datalaria_EN.xlsx")

    # 1. Versión en Español
    wb_es = build_workbook(lang='ES')
    wb_es.save(path_es)
    print(f"-> Guardado exitosamente: {path_es}")

    # 2. Versión en Inglés
    wb_en = build_workbook(lang='EN')
    wb_en.save(path_en)
    print(f"-> Guardado exitosamente: {path_en}")

    print("\nGeneración de modelos Excel EVM completada con éxito.")


if __name__ == "__main__":
    generate_all_excel()
