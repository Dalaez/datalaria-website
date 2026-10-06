import openpyxl

def audit_workbook(path):
    print(f"=== AUDITING {path} ===")
    wb = openpyxl.load_workbook(path, data_only=False)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        prot = ws.protection
        print(f"\n--- Sheet: {sheetname} ---")
        print(f"Protection: sheet={prot.sheet}, selectUnlockedCells={prot.selectUnlockedCells}, selectLockedCells={prot.selectLockedCells}")
        
        formula_cells = []
        unlocked_cells = []
        locked_no_formula = []
        unlocked_with_formula = []

        for row in ws.rows:
            for cell in row:
                if cell.value is None:
                    continue
                val_str = str(cell.value)
                is_formula = val_str.startswith('=')
                is_locked = cell.protection.locked if cell.protection else True

                if is_formula:
                    formula_cells.append(cell.coordinate)
                    if not is_locked:
                        unlocked_with_formula.append(cell.coordinate)
                else:
                    if not is_locked:
                        unlocked_cells.append((cell.coordinate, val_str[:25]))
                    else:
                        locked_no_formula.append(cell.coordinate)

        print(f"  Formulas count: {len(formula_cells)} (Unlocked formulas: {len(unlocked_with_formula)})")
        print(f"  Unlocked input cells: {len(unlocked_cells)}")
        print(f"  Locked header/static cells: {len(locked_no_formula)}")

        # Check Veto sheet row 8
        if "Veto" in sheetname:
            for coord in ['B8', 'C8', 'D8', 'E8', 'F8', 'G8', 'H8', 'I8', 'J8']:
                c = ws[coord]
                col_rgb = c.fill.start_color.rgb if c.fill else None
                print(f"    Veto Cell {coord}: val='{str(c.value)[:25]}', locked={c.protection.locked}, fill={col_rgb}")

audit_workbook("packages/Suite_02_Toma_Decisiones/01_Matriz_DAR_Cuantitativa/[ES]_Matriz_DAR/Matriz_DAR_Cuantitativa_Datalaria_ES.xlsx")
print("\n" + "="*60 + "\n")
audit_workbook("packages/Suite_02_Toma_Decisiones/01_Matriz_DAR_Cuantitativa/[EN]_Quantitative_DAR/Quantitative_DAR_Matrix_Datalaria_EN.xlsx")
